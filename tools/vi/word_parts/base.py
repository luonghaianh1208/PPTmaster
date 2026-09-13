"""Tầng dựng file Word dùng chung: khổ giấy, font, bảng, chân trang.

Gói nào của lớp Việt cũng dùng module này; phần bố cục riêng thì để ở gói đó.
"""

from __future__ import annotations

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from . import inline


def field(paragraph, instruction: str) -> None:
    """Chèn một trường Word (PAGE, NUMPAGES) vào đoạn."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction_element = OxmlElement("w:instrText")
    instruction_element.set(qn("xml:space"), "preserve")
    instruction_element.text = f" {instruction} "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instruction_element)
    run._r.append(end)


def add_page_numbers(section) -> None:
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run("Trang ")
    field(paragraph, "PAGE")
    paragraph.add_run(" / ")
    field(paragraph, "NUMPAGES")


def new_document(
    *,
    width_cm: float = 21.0,
    height_cm: float = 29.7,
    margins_cm: tuple[float, float, float, float] = (1.8, 1.8, 2.5, 1.5),
    font: str = "Times New Roman",
    size_pt: int = 12,
    line_spacing: float = 1.15,
    space_after_pt: int = 2,
    page_numbers: bool = True,
) -> Document:
    """margins_cm theo thứ tự trên, dưới, trái, phải."""
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(width_cm)
    section.page_height = Cm(height_cm)
    top, bottom, left, right = margins_cm
    section.top_margin = Cm(top)
    section.bottom_margin = Cm(bottom)
    section.left_margin = Cm(left)
    section.right_margin = Cm(right)
    normal = document.styles["Normal"]
    normal.font.name = font
    normal.font.size = Pt(size_pt)
    fonts = normal.element.get_or_add_rPr().get_or_add_rFonts()
    for attribute in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attribute), font)
    normal.paragraph_format.line_spacing = line_spacing
    normal.paragraph_format.space_after = Pt(space_after_pt)
    if page_numbers:
        add_page_numbers(section)
    return document


def write(paragraph, text: str, *, italic: bool = False, bold: bool = False, color=None) -> None:
    """Ghi text vào đoạn, dịch ~chỉ số dưới~, ^chỉ số trên^ và **in đậm** thành run thật."""
    for chunk, kind in inline.split_runs(text):
        run = paragraph.add_run(chunk)
        run.italic = italic
        run.bold = bold or kind == inline.BOLD
        if kind == inline.SUB:
            run.font.subscript = True
        elif kind == inline.SUP:
            run.font.superscript = True
        if color is not None:
            run.font.color.rgb = color


def _borders(table, value: str, size: str) -> None:
    tbl_pr = table._tbl.tblPr
    for existing in tbl_pr.findall(qn("w:tblBorders")):
        tbl_pr.remove(existing)
    element = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge_element = OxmlElement(f"w:{edge}")
        edge_element.set(qn("w:val"), value)
        edge_element.set(qn("w:sz"), size)
        element.append(edge_element)
    tbl_pr.insert_element_before(
        element,
        "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
        "w:tblCaption", "w:tblDescription", "w:tblPrChange",
    )


def clear_borders(table) -> None:
    _borders(table, "none", "0")


def grid_borders(table) -> None:
    _borders(table, "single", "6")


def set_widths(table, widths) -> None:
    table.autofit = False
    for index, width in enumerate(widths):
        table.columns[index].width = width
        for cell in table.columns[index].cells:
            cell.width = width


def cell_text(cell, text: str, *, bold: bool = False) -> None:
    cell.paragraphs[0].text = ""
    write(cell.paragraphs[0], text, bold=bold)


def fill_cell(cell, lines, *, bold_first: bool = True) -> None:
    """Mỗi dòng thành một đoạn căn giữa trong ô; dòng đầu có thể in đậm."""
    cell.paragraphs[0].text = ""
    for index, line in enumerate(lines):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        write(paragraph, line, bold=bold_first and index == 0)
