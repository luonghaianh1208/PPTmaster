"""Dựng file Word Kế hoạch bài dạy theo Công văn 5512, và file ghi chú nội bộ.

Thể thức: A4 dọc, lề 2/2/2,5/2 cm, Times New Roman 14pt, giãn dòng 1,3.
Ghi chú nội bộ không bao giờ vào file giáo án; nó ra can-soat.md.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from word_parts import base
from word_parts.base import cell_text, clear_borders, fill_cell, grid_borders, set_widths, write

from .frameworks import TRACK_ALIASES, normalise
from .parse import Activity, Lesson

_ACTIVITY_PREFIX_RE = re.compile(r"^hoạt\s*động\s+\d+\s*[:.\-–—]?\s*", re.IGNORECASE)

FILENAME = "giao-an.docx"
REVIEW_FILENAME = "can-soat.md"
MARGINS = (2.0, 2.0, 2.5, 2.0)
FONT_SIZE = 14
TABLE_SIZE = 12
LINE_SPACING = 1.3
NO_REVIEW = "Không có mục nào cần soát."
STEPS = (
    ("chuyen_giao", "Bước 1. Chuyển giao nhiệm vụ"),
    ("thuc_hien", "Bước 2. Học sinh thực hiện nhiệm vụ"),
    ("bao_cao", "Bước 3. Báo cáo, thảo luận"),
    ("ket_luan", "Bước 4. Kết luận, nhận định"),
)
OBJECTIVE_GROUPS = (
    ("a) Năng lực chung", "general"),
    ("b) Năng lực đặc thù", "specific"),
    ("c) Năng lực số", "nls_lines"),
    ("d) Năng lực trí tuệ nhân tạo (AI)", "ai_lines"),
)


def _document() -> Document:
    return base.new_document(
        margins_cm=MARGINS,
        size_pt=FONT_SIZE,
        line_spacing=LINE_SPACING,
        space_after_pt=4,
    )


def _paragraph(document: Document, text: str, alignment, *, bold: bool = False,
               italic: bool = False, space_before: int = 0):
    paragraph = document.add_paragraph()
    paragraph.alignment = alignment
    if space_before:
        paragraph.paragraph_format.space_before = Pt(space_before)
    write(paragraph, text, bold=bold, italic=italic)
    return paragraph


def _justified(document: Document, text: str, **kwargs):
    return _paragraph(document, text, WD_ALIGN_PARAGRAPH.JUSTIFY, **kwargs)


def _left(document: Document, text: str, **kwargs):
    return _paragraph(document, text, WD_ALIGN_PARAGRAPH.LEFT, **kwargs)


def _centre(document: Document, text: str, **kwargs):
    return _paragraph(document, text, WD_ALIGN_PARAGRAPH.CENTER, **kwargs)


def _shrink(table) -> None:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(TABLE_SIZE)


def _add_header(document: Document, lesson: Lesson) -> None:
    table = document.add_table(rows=1, cols=2)
    clear_borders(table)
    set_widths(table, [Cm(8.0), Cm(8.5)])
    left_lines = [lesson.meta["school"]]
    if lesson.meta.get("department"):
        left_lines.append(lesson.meta["department"])
    fill_cell(table.rows[0].cells[0], left_lines)
    right_lines = []
    if lesson.meta.get("teacher"):
        right_lines.append(f"Giáo viên: {lesson.meta['teacher']}")
    if lesson.meta.get("school_year"):
        right_lines.append(f"Năm học: {lesson.meta['school_year']}")
    fill_cell(table.rows[0].cells[1], right_lines or [""], bold_first=False)
    _shrink(table)


def _add_title(document: Document, lesson: Lesson) -> None:
    _centre(document, "KẾ HOẠCH BÀI DẠY", bold=True, space_before=8)
    _centre(document, lesson.meta["lesson"], bold=True)
    detail = f"Môn: {lesson.meta['subject']} — Lớp {lesson.meta['grade']}"
    if normalise(lesson.meta.get("track", "")) in TRACK_ALIASES:
        detail += " (hệ chuyên)"
    _centre(document, detail)
    schedule = f"Thời lượng: {lesson.meta['periods']} tiết"
    if lesson.meta.get("period_numbers"):
        schedule += f" — Tiết thứ {lesson.meta['period_numbers']}"
    if lesson.meta.get("week"):
        schedule += f" — Tuần {lesson.meta['week']}"
    _centre(document, schedule)


def _add_objectives(document: Document, lesson: Lesson) -> None:
    _justified(document, "I. MỤC TIÊU", bold=True, space_before=8)
    _justified(document, "1. Về kiến thức", bold=True)
    for item in lesson.knowledge:
        _justified(document, f"- {item}")
    _justified(document, "2. Về năng lực", bold=True)
    for label, attribute in OBJECTIVE_GROUPS:
        _justified(document, label, italic=True)
        for item in getattr(lesson, attribute):
            _justified(document, f"- {item}")
    _justified(document, "3. Về phẩm chất", bold=True)
    for item in lesson.qualities:
        _justified(document, f"- {item}")


def _add_equipment(document: Document, lesson: Lesson) -> None:
    _justified(document, "II. THIẾT BỊ DẠY HỌC VÀ HỌC LIỆU", bold=True, space_before=8)
    _justified(document, "1. Chuẩn bị của giáo viên", bold=True)
    for item in lesson.teacher_tools:
        _justified(document, f"- {item}")
    _justified(document, "2. Chuẩn bị của học sinh", bold=True)
    for item in lesson.student_tools:
        _justified(document, f"- {item}")


def _add_activity(document: Document, activity: Activity, number: int) -> None:
    title = _ACTIVITY_PREFIX_RE.sub("", activity.title)
    _justified(
        document,
        f"Hoạt động {number}. {title} ({activity.minutes} phút)",
        bold=True,
        space_before=8,
    )
    _justified(document, "a) Mục tiêu", italic=True)
    for text in activity.muc_tieu:
        _justified(document, text)
    if activity.nls:
        _justified(document, "Năng lực số: " + ", ".join(activity.nls))
    if activity.ai:
        _justified(document, "Năng lực AI: " + ", ".join(activity.ai))
    _justified(document, "b) Nội dung", italic=True)
    for text in activity.noi_dung:
        _justified(document, text)
    _justified(document, "c) Sản phẩm", italic=True)
    for text in activity.san_pham:
        _justified(document, text)
    _justified(document, "d) Tổ chức thực hiện", italic=True)
    for attribute, label in STEPS:
        _justified(document, label, bold=True)
        for text in getattr(activity, attribute):
            _justified(document, text)


def _add_appendix(document: Document, lesson: Lesson) -> None:
    _justified(document, "IV. PHỤ LỤC", bold=True, space_before=8)
    for sheet in lesson.worksheets:
        _left(document, sheet.title, bold=True, space_before=6)
        for item in sheet.items:
            _left(document, item)
    _justified(document, "Rubric đánh giá năng lực số và năng lực AI", bold=True, space_before=8)
    table = document.add_table(rows=len(lesson.rubrics) + 1, cols=4)
    grid_borders(table)
    set_widths(table, [Cm(4.5), Cm(4.0), Cm(4.0), Cm(4.0)])
    for cell, text in zip(table.rows[0].cells, ("Tiêu chí", "Mức 1", "Mức 2", "Mức 3")):
        cell_text(cell, text, bold=True)
    for row_index, rubric in enumerate(lesson.rubrics, start=1):
        cells = table.rows[row_index].cells
        cell_text(cells[0], rubric.title)
        for column, level in enumerate(rubric.levels, start=1):
            cell_text(cells[column], level)
    _shrink(table)


def build(lesson: Lesson, folder: Path) -> Path:
    document = _document()
    _add_header(document, lesson)
    _add_title(document, lesson)
    _add_objectives(document, lesson)
    _add_equipment(document, lesson)
    _justified(document, "III. TIẾN TRÌNH DẠY HỌC", bold=True, space_before=8)
    for number, activity in enumerate(lesson.activities, start=1):
        _add_activity(document, activity, number)
    _add_appendix(document, lesson)
    path = folder / FILENAME
    document.save(str(path))
    return path


def write_review(lesson: Lesson, folder: Path, warnings: list[str]) -> Path:
    path = folder / REVIEW_FILENAME
    lines = [
        "# Cần thầy cô soát",
        "",
        f"Bài: {lesson.meta['lesson']}",
        f"Môn: {lesson.meta['subject']} — Lớp {lesson.meta['grade']}",
        "",
        "File này là ghi chú nội bộ, không nằm trong giáo án nộp cho trường.",
        "",
    ]
    items = list(lesson.review_notes) + list(warnings)
    lines += [f"- {item}" for item in items] if items else [NO_REVIEW]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path
