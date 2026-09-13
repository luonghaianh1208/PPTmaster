"""Test cho lớp soạn đề KHTN tiếng Anh của bản Việt."""

import contextlib
import io
import json
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest import mock

REPO_ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(REPO_ROOT / "tools" / "vi"))

from de_thi_parts import docx_build, parse  # noqa: E402
from word_parts import inline  # noqa: E402

import de_thi  # noqa: E402


class InlineTest(unittest.TestCase):
    def test_subscript_marks_become_sub_runs(self):
        self.assertEqual(
            inline.split_runs("H~2~SO~4~"),
            [("H", inline.PLAIN), ("2", inline.SUB), ("SO", inline.PLAIN), ("4", inline.SUB)],
        )

    def test_superscript_marks_become_sup_runs(self):
        self.assertEqual(
            inline.split_runs("5.0 m/s^2^"),
            [("5.0 m/s", inline.PLAIN), ("2", inline.SUP)],
        )

    def test_bold_marks_become_bold_runs(self):
        self.assertEqual(
            inline.split_runs("Which one is **not** correct?"),
            [("Which one is ", inline.PLAIN), ("not", inline.BOLD), (" correct?", inline.PLAIN)],
        )

    def test_plain_text_has_one_run(self):
        self.assertEqual(inline.split_runs("no marks here"), [("no marks here", inline.PLAIN)])

    def test_empty_text_still_returns_one_run(self):
        self.assertEqual(inline.split_runs(""), [("", inline.PLAIN)])

    def test_unpaired_marks_stay_plain(self):
        self.assertEqual(inline.split_runs("a ~ b ^ c"), [("a ~ b ^ c", inline.PLAIN)])

    def test_plain_text_strips_every_mark(self):
        self.assertEqual(inline.plain_text("H~2~O at 5 m/s^2^ is **not** ice"), "H2O at 5 m/s2 is not ice")


VALID_SOURCE = """---
school: TRƯỜNG THPT CHUYÊN NGUYỄN TRÃI
department: TỔ VẬT LÍ
title: ĐỀ KIỂM TRA GIỮA HỌC KÌ I
subject: PHYSICS — Grade 10
time: 45
code: 101
---

## PART I

### 1
en: An object of mass 2.0 kg is acted on by a resultant force of 10 N. Find its acceleration.
vi: Một vật khối lượng 2,0 kg chịu tác dụng của lực tổng hợp 10 N. Tính gia tốc của vật.
A: 2.0 m/s^2^
B: 5.0 m/s^2^
C: 10 m/s^2^
D: 20 m/s^2^
key: B
why: a = F/m = 10/2.0 = 5.0 m/s^2^
level: hieu
topic: Dynamics

### 2
en: Which of the following is **not** a vector quantity?
vi: Đại lượng nào sau đây **không** phải đại lượng vectơ?
A: Velocity
B: Force
C: Speed
D: Acceleration
key: C
level: biet
topic: Kinematics

## PART II

### 1
en: A trolley moves down a smooth slope.
vi: Một xe lăn đi xuống mặt phẳng nghiêng nhẵn.
a: Its speed increases. | T
b: Its acceleration is zero. | F
c: The resultant force acts down the slope. | T
d: Its kinetic energy decreases. | F
level: vandung
topic: Dynamics

## PART III

### 1
en: Calculate the work done by a force of 12 N moving an object 3.0 m in the direction of the force.
vi: Tính công của lực 12 N khi vật dịch chuyển 3,0 m theo hướng của lực.
key: 36
unit: J
level: hieu
topic: Energy

## CAN SOAT
- Thuật ngữ "resultant force" đã dùng thống nhất, thầy cô xác nhận lại giúp em.
"""


def source_without(line_prefix: str) -> str:
    kept = [line for line in VALID_SOURCE.splitlines() if not line.startswith(line_prefix)]
    return "\n".join(kept) + "\n"


class ParseTest(unittest.TestCase):
    def test_valid_source_reads_every_part(self):
        exam = parse.parse_exam(VALID_SOURCE)
        self.assertEqual(exam.counts(), {"part1": 2, "part2": 1, "part3": 1})
        self.assertEqual(exam.meta["school"], "TRƯỜNG THPT CHUYÊN NGUYỄN TRÃI")
        self.assertEqual(exam.meta["time"], "45")
        self.assertEqual(exam.meta["code"], "101")

    def test_multiple_choice_question_keeps_options_and_key(self):
        question = parse.parse_exam(VALID_SOURCE).part(1)[0]
        self.assertEqual(sorted(question.options), ["A", "B", "C", "D"])
        self.assertEqual(question.key, "B")
        self.assertEqual(question.level, "hieu")
        self.assertEqual(question.topic, "Dynamics")
        self.assertTrue(question.why)

    def test_true_false_question_keeps_four_statements_with_flags(self):
        question = parse.parse_exam(VALID_SOURCE).part(2)[0]
        self.assertEqual([label for label, _, _ in question.statements], ["a", "b", "c", "d"])
        self.assertEqual([correct for _, _, correct in question.statements], [True, False, True, False])

    def test_short_answer_question_keeps_key_and_unit(self):
        question = parse.parse_exam(VALID_SOURCE).part(3)[0]
        self.assertEqual(question.key, "36")
        self.assertEqual(question.unit, "J")

    def test_review_notes_are_collected(self):
        exam = parse.parse_exam(VALID_SOURCE)
        self.assertEqual(len(exam.review_notes), 1)
        self.assertIn("resultant force", exam.review_notes[0])

    def test_missing_vietnamese_line_is_a_warning_not_an_error(self):
        exam = parse.parse_exam(source_without("vi: Một vật khối lượng"))
        self.assertTrue(any("chưa có bản tiếng Việt" in warning for warning in exam.warnings))
        self.assertEqual(exam.part(1)[0].vi, "")

    def test_file_without_meta_block_fails_on_line_one(self):
        with self.assertRaises(parse.ParseError) as caught:
            parse.parse_exam("## PART I\n")
        self.assertEqual(caught.exception.line_no, 1)

    def test_meta_missing_required_key_names_it(self):
        with self.assertRaises(parse.ParseError) as caught:
            parse.parse_exam(source_without("subject:"))
        self.assertIn("subject", caught.exception.message)

    def test_meta_time_must_be_digits(self):
        broken = VALID_SOURCE.replace("time: 45", "time: bốn mươi lăm")
        with self.assertRaises(parse.ParseError) as caught:
            parse.parse_exam(broken)
        self.assertIn("time", caught.exception.message)

    def test_meta_rejects_unknown_key(self):
        broken = VALID_SOURCE.replace("code: 101", "mon: Vật lí")
        with self.assertRaises(parse.ParseError) as caught:
            parse.parse_exam(broken)
        self.assertIn("mon", caught.exception.message)

    def test_missing_option_names_the_label_and_line(self):
        with self.assertRaises(parse.ParseError) as caught:
            parse.parse_exam(source_without("C: 10 m/s^2^"))
        self.assertIn("'C'", caught.exception.message)
        self.assertGreater(caught.exception.line_no, 1)

    def test_key_outside_abcd_fails(self):
        broken = VALID_SOURCE.replace("key: B", "key: E", 1)
        with self.assertRaises(parse.ParseError) as caught:
            parse.parse_exam(broken)
        self.assertIn("A, B, C, D", caught.exception.message)

    def test_missing_key_in_part_one_fails(self):
        with self.assertRaises(parse.ParseError) as caught:
            parse.parse_exam(source_without("key: B"))
        self.assertIn("key", caught.exception.message)

    def test_level_must_be_one_of_three_values(self):
        broken = VALID_SOURCE.replace("level: hieu", "level: kho", 1)
        with self.assertRaises(parse.ParseError) as caught:
            parse.parse_exam(broken)
        for value in parse.LEVELS:
            self.assertIn(value, caught.exception.message)

    def test_true_false_statement_needs_a_flag(self):
        broken = VALID_SOURCE.replace("a: Its speed increases. | T", "a: Its speed increases.")
        with self.assertRaises(parse.ParseError) as caught:
            parse.parse_exam(broken)
        self.assertIn("' | T'", caught.exception.message)

    def test_parts_must_be_in_order(self):
        broken = VALID_SOURCE.replace("## PART I\n", "## PART III\n", 1)
        with self.assertRaises(parse.ParseError) as caught:
            parse.parse_exam(broken)
        self.assertIn("thứ tự", caught.exception.message)

    def test_question_numbers_must_restart_from_one_in_each_part(self):
        broken = VALID_SOURCE.replace("### 2\n", "### 5\n", 1)
        with self.assertRaises(parse.ParseError) as caught:
            parse.parse_exam(broken)
        self.assertIn("2", caught.exception.message)

    def test_exam_with_no_question_fails(self):
        header = VALID_SOURCE.split("## PART I")[0]
        with self.assertRaises(parse.ParseError) as caught:
            parse.parse_exam(header)
        self.assertIn("không có câu hỏi", caught.exception.message)

    def test_duplicate_key_in_one_question_fails(self):
        broken = VALID_SOURCE.replace("level: hieu\ntopic: Dynamics", "level: hieu\nlevel: biet", 1)
        with self.assertRaises(parse.ParseError) as caught:
            parse.parse_exam(broken)
        self.assertIn("hai lần", caught.exception.message)

    def test_review_section_must_be_last(self):
        broken = VALID_SOURCE.replace("## PART II", "## CAN SOAT\n- ghi chú\n\n## PART II", 1)
        with self.assertRaises(parse.ParseError) as caught:
            parse.parse_exam(broken)
        self.assertIn("mục cuối", caught.exception.message)

    def test_unit_key_is_rejected_in_part_one(self):
        broken = VALID_SOURCE.replace("key: B", "key: B\nunit: J", 1)
        with self.assertRaises(parse.ParseError) as caught:
            parse.parse_exam(broken)
        self.assertIn("unit", caught.exception.message)

    def test_superscript_digit_in_question_number_raises_parse_error(self):
        broken = VALID_SOURCE.replace("### 1\n", "### ²\n", 1)
        with self.assertRaises(parse.ParseError) as caught:
            parse.parse_exam(broken)
        self.assertNotIsInstance(caught.exception, ValueError)
        expected_line = broken.splitlines().index("### ²") + 1
        self.assertEqual(caught.exception.line_no, expected_line)

    def test_superscript_digit_in_time_raises_parse_error(self):
        broken = VALID_SOURCE.replace("time: 45", "time: ²")
        with self.assertRaises(parse.ParseError) as caught:
            parse.parse_exam(broken)
        self.assertIn("time", caught.exception.message)

    def test_part_three_key_must_be_numeric(self):
        broken = VALID_SOURCE.replace("key: 36", "key: ba mươi sáu")
        with self.assertRaises(parse.ParseError) as caught:
            parse.parse_exam(broken)
        self.assertIn("PART III", caught.exception.message)
        expected_line = broken.splitlines().index("key: ba mươi sáu") + 1
        self.assertEqual(caught.exception.line_no, expected_line)

    def test_part_three_key_accepts_decimal_and_negative_numbers(self):
        source_with_decimal = VALID_SOURCE.replace("key: 36", "key: 12,5")
        exam1 = parse.parse_exam(source_with_decimal)
        self.assertEqual(exam1.part(3)[0].key, "12,5")
        source_with_negative = VALID_SOURCE.replace("key: 36", "key: -0.25")
        exam2 = parse.parse_exam(source_with_negative)
        self.assertEqual(exam2.part(3)[0].key, "-0.25")


class ExamMathTest(unittest.TestCase):
    def build(self, part1: int, part2: int, part3: int, points: str = "10") -> parse.Exam:
        lines = [
            "---",
            "school: TRƯỜNG A",
            "title: ĐỀ THI",
            "subject: PHYSICS — Grade 10",
            "time: 50",
            f"points: {points}",
            "---",
        ]
        if part1:
            lines.append("## PART I")
            for number in range(1, part1 + 1):
                lines += [f"### {number}", "en: Q", "A: 1", "B: 2", "C: 3", "D: 4",
                          "key: A", "level: biet", "topic: Motion"]
        if part2:
            lines.append("## PART II")
            for number in range(1, part2 + 1):
                lines += [f"### {number}", "en: Q",
                          "a: s | T", "b: s | F", "c: s | T", "d: s | F",
                          "level: hieu", "topic: Energy"]
        if part3:
            lines.append("## PART III")
            for number in range(1, part3 + 1):
                lines += [f"### {number}", "en: Q", "key: 1", "level: vandung"]
        return parse.parse_exam("\n".join(lines) + "\n")

    def test_standard_paper_totals_ten_points(self):
        exam = self.build(18, 4, 6)
        self.assertEqual(exam.total_points(), 10.0)
        self.assertEqual(exam.points_warning(), "")

    def test_wrong_count_produces_a_points_warning(self):
        exam = self.build(10, 4, 6)
        self.assertIn("points", exam.points_warning())

    def test_matrix_groups_by_topic_and_level(self):
        exam = self.build(2, 1, 0)
        rows = dict((row[0], row[1:]) for row in exam.matrix())
        self.assertEqual(rows["Motion"], (2, 0, 0, 2))
        self.assertEqual(rows["Energy"], (0, 1, 0, 1))

    def test_questions_without_topic_land_in_one_row(self):
        exam = self.build(0, 0, 2)
        self.assertEqual(exam.matrix(), [(parse.NO_TOPIC, 0, 0, 2, 2)])


def document_xml(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        return archive.read("word/document.xml").decode("utf-8")


def footer_xml(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        names = [name for name in archive.namelist() if name.startswith("word/footer")]
        if not names:
            return ""
        return archive.read(names[0]).decode("utf-8")


class OptionLayoutTest(unittest.TestCase):
    def test_short_options_use_four_columns(self):
        options = {"A": "2.0 m/s^2^", "B": "5.0 m/s^2^", "C": "10 m/s^2^", "D": "20 m/s^2^"}
        self.assertEqual(docx_build.option_columns(options), 4)

    def test_medium_options_use_two_columns(self):
        options = {"A": "kinetic friction", "B": "static friction", "C": "air resistance", "D": "normal contact force"}
        self.assertEqual(docx_build.option_columns(options), 2)

    def test_long_options_use_one_column(self):
        long_text = "The resultant force acting on the trolley points down the slope at all times"
        options = {"A": long_text, "B": long_text, "C": long_text, "D": long_text}
        self.assertEqual(docx_build.option_columns(options), 1)

    def test_marks_do_not_count_towards_option_length(self):
        options = {"A": "H~2~SO~4~", "B": "HCl", "C": "NaOH", "D": "KOH"}
        self.assertEqual(docx_build.option_columns(options), 4)


class DocxBuildTest(unittest.TestCase):
    def setUp(self):
        self.exam = parse.parse_exam(VALID_SOURCE)
        self.tmp = tempfile.TemporaryDirectory()
        self.folder = Path(self.tmp.name)
        self.addCleanup(self.tmp.cleanup)

    def test_paper_uses_a4_and_exam_margins(self):
        from docx import Document

        path = docx_build.build_de(self.exam, self.folder / "de-en.docx")
        section = Document(str(path)).sections[0]
        # Word lưu khổ giấy theo twip nên đọc lại lệch vài trăm EMU (< 0,001 cm); so tới 0,01 cm.
        for attribute, expected_cm in (
            ("page_width", 21.0),
            ("page_height", 29.7),
            ("top_margin", 1.8),
            ("bottom_margin", 1.8),
            ("left_margin", 2.5),
            ("right_margin", 1.5),
        ):
            with self.subTest(attribute=attribute):
                self.assertAlmostEqual(getattr(section, attribute).cm, expected_cm, places=2)

    def test_body_font_is_times_new_roman_twelve(self):
        from docx import Document
        from docx.shared import Pt

        path = docx_build.build_de(self.exam, self.folder / "de-en.docx")
        normal = Document(str(path)).styles["Normal"]
        self.assertEqual(normal.font.name, "Times New Roman")
        self.assertEqual(normal.font.size, Pt(12))

    def test_footer_carries_page_number_fields(self):
        path = docx_build.build_de(self.exam, self.folder / "de-en.docx")
        xml = footer_xml(path)
        self.assertIn("PAGE", xml)
        self.assertIn("NUMPAGES", xml)
        self.assertIn("Trang", xml)

    def test_subscript_marks_become_real_subscript(self):
        source = VALID_SOURCE.replace("A: 2.0 m/s^2^", "A: H~2~SO~4~")
        exam = parse.parse_exam(source)
        path = docx_build.build_de(exam, self.folder / "de-en.docx")
        self.assertIn('w:val="subscript"', document_xml(path))

    def test_superscript_marks_become_real_superscript(self):
        path = docx_build.build_de(self.exam, self.folder / "de-en.docx")
        self.assertIn('w:val="superscript"', document_xml(path))

    def test_student_paper_hides_answers_and_explanations(self):
        source = VALID_SOURCE.replace(
            "why: a = F/m = 10/2.0 = 5.0 m/s^2^", "why: DAU-HIEU-GIAI-THICH"
        )
        exam = parse.parse_exam(source)
        path = docx_build.build_de(exam, self.folder / "de-en.docx")
        xml = document_xml(path)
        self.assertNotIn("DAU-HIEU-GIAI-THICH", xml)
        self.assertNotIn("Đáp án", xml)

    def test_student_paper_has_the_candidate_line_and_end_marker(self):
        path = docx_build.build_de(self.exam, self.folder / "de-en.docx")
        xml = document_xml(path)
        self.assertIn("Full name", xml)
        self.assertIn("THE END", xml)

    def test_bilingual_paper_carries_the_vietnamese_line(self):
        path = docx_build.build_de(self.exam, self.folder / "song-ngu.docx", bilingual=True)
        xml = document_xml(path)
        self.assertIn("Tính gia tốc của vật", xml)

    def test_student_paper_has_no_vietnamese_question_text(self):
        path = docx_build.build_de(self.exam, self.folder / "de-en.docx")
        self.assertNotIn("Tính gia tốc của vật", document_xml(path))

    def test_answer_file_has_keys_scoring_matrix_and_review(self):
        path = docx_build.build_dap_an(self.exam, self.folder / "dap-an.docx", [])
        xml = document_xml(path)
        for expected in ("Thang điểm", "Ma trận đặc tả", "Cần thầy cô soát", "Hướng dẫn giải", "Dynamics"):
            self.assertIn(expected, xml)

    def test_answer_file_repeats_tool_warnings_on_paper(self):
        path = docx_build.build_dap_an(self.exam, self.folder / "dap-an.docx", ["CANH-BAO-THU-NGHIEM"])
        self.assertIn("CANH-BAO-THU-NGHIEM", document_xml(path))

    def test_answer_file_says_so_when_nothing_needs_review(self):
        source = VALID_SOURCE.split("## CAN SOAT")[0]
        exam = parse.parse_exam(source)
        path = docx_build.build_dap_an(exam, self.folder / "dap-an.docx", [])
        self.assertIn("Không có mục nào cần soát", document_xml(path))

    def test_build_writes_every_requested_file(self):
        written = docx_build.build(self.exam, self.folder, ["de", "song-ngu", "dap-an"], [])
        self.assertEqual([path.name for path in written],
                         ["de-en.docx", "de-song-ngu.docx", "dap-an.docx"])
        for path in written:
            self.assertTrue(path.is_file(), path)

    def test_build_writes_only_the_requested_subset(self):
        written = docx_build.build(self.exam, self.folder, ["de", "dap-an"], [])
        self.assertEqual([path.name for path in written], ["de-en.docx", "dap-an.docx"])
        self.assertFalse((self.folder / "de-song-ngu.docx").exists())

    @staticmethod
    def border_values(table):
        from docx.oxml.ns import qn

        borders = table._tbl.tblPr.find(qn("w:tblBorders"))
        if borders is None:
            return set()
        return {edge.get(qn("w:val")) for edge in borders}

    @staticmethod
    def option_tables(path):
        from docx import Document

        return [table for table in Document(str(path)).tables
                if table.rows[0].cells[0].text.startswith("A.")]

    def test_header_is_borderless_and_true_false_table_is_a_grid(self):
        from docx import Document

        path = docx_build.build_de(self.exam, self.folder / "de-en.docx")
        tables = Document(str(path)).tables
        self.assertEqual(self.border_values(tables[0]), {"none"})
        statements = [table for table in tables if table.rows[0].cells[0].text == "Statement"]
        self.assertEqual(len(statements), 1)
        self.assertEqual(self.border_values(statements[0]), {"single"})

    def test_short_options_render_as_one_row_of_four_columns(self):
        path = docx_build.build_de(self.exam, self.folder / "de-en.docx")
        tables = self.option_tables(path)
        self.assertEqual(len(tables), 2)
        for table in tables:
            self.assertEqual((len(table.rows), len(table.columns)), (1, 4))
            self.assertEqual(self.border_values(table), {"none"})

    def test_long_options_render_one_per_row(self):
        long_text = "The resultant force acting on the trolley points down the slope"
        source = VALID_SOURCE
        for label, old in (("A", "2.0 m/s^2^"), ("B", "5.0 m/s^2^"), ("C", "10 m/s^2^"), ("D", "20 m/s^2^")):
            source = source.replace(f"{label}: {old}", f"{label}: {long_text} {label}", 1)
        exam = parse.parse_exam(source)
        path = docx_build.build_de(exam, self.folder / "de-en.docx")
        first = self.option_tables(path)[0]
        self.assertEqual((len(first.rows), len(first.columns)), (4, 1))


class SelectPartsTest(unittest.TestCase):
    def test_default_selects_every_part(self):
        self.assertEqual(de_thi.select_parts("tat-ca"), ["de", "song-ngu", "dap-an"])

    def test_comma_list_keeps_the_given_order(self):
        self.assertEqual(de_thi.select_parts("dap-an,de"), ["dap-an", "de"])

    def test_duplicates_are_dropped(self):
        self.assertEqual(de_thi.select_parts("de,de"), ["de"])

    def test_unknown_value_raises(self):
        with self.assertRaises(ValueError):
            de_thi.select_parts("dap_an")

    def test_empty_value_raises(self):
        with self.assertRaises(ValueError):
            de_thi.select_parts("")


class CliTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.folder = Path(self.tmp.name) / "đề thi vật lí 10"
        self.folder.mkdir(parents=True)
        self.addCleanup(self.tmp.cleanup)

    def write_source(self, text: str = VALID_SOURCE) -> None:
        (self.folder / "de.md").write_text(text, encoding="utf-8")

    def run_cli(self, *args: str) -> tuple[int, dict]:
        out = io.StringIO()
        with contextlib.redirect_stdout(out), contextlib.redirect_stderr(io.StringIO()):
            code = de_thi.main([str(self.folder), *args])
        printed = out.getvalue().strip().splitlines()
        self.assertEqual(len(printed), 1, printed)
        return code, json.loads(printed[0])

    def test_full_run_writes_three_files(self):
        self.write_source()
        code, data = self.run_cli()
        self.assertEqual(code, 0, data)
        self.assertTrue(data["ready"])
        self.assertEqual(len(data["files"]), 3)
        self.assertEqual(data["questions"], {"part1": 2, "part2": 1, "part3": 1})
        self.assertIsNone(data["error"])

    def test_run_accepts_a_vietnamese_folder_name(self):
        self.write_source()
        code, data = self.run_cli()
        self.assertEqual(code, 0, data)
        for name in data["files"]:
            self.assertTrue(Path(name).is_file(), name)

    def test_subset_writes_only_two_files(self):
        self.write_source()
        code, data = self.run_cli("--phan", "de,dap-an")
        self.assertEqual(code, 0, data)
        self.assertEqual([Path(name).name for name in data["files"]], ["de-en.docx", "dap-an.docx"])

    def test_plan_only_writes_nothing(self):
        self.write_source()
        code, data = self.run_cli("--plan-only")
        self.assertEqual(code, 0, data)
        self.assertEqual(data["files"], [])
        self.assertEqual(list(self.folder.glob("*.docx")), [])

    def test_missing_source_reports_input_step(self):
        code, data = self.run_cli()
        self.assertEqual(code, 1)
        self.assertEqual(data["error"]["step"], "input")
        self.assertIn("de.md", data["error"]["message"])

    def test_broken_source_reports_parse_step_with_the_line(self):
        self.write_source(VALID_SOURCE.replace("key: B", "key: E", 1))
        code, data = self.run_cli()
        self.assertEqual(code, 1)
        self.assertEqual(data["error"]["step"], "parse")
        self.assertIn("Dòng", data["error"]["message"])

    def test_unknown_part_value_reports_input_step(self):
        self.write_source()
        code, data = self.run_cli("--phan", "dap_an")
        self.assertEqual(code, 1)
        self.assertEqual(data["error"]["step"], "input")

    def test_missing_python_docx_reports_docx_step_with_the_install_command(self):
        self.write_source()
        with mock.patch.object(de_thi, "load_docx_build", side_effect=ImportError("No module named 'docx'")):
            code, data = self.run_cli()
        self.assertEqual(code, 1)
        self.assertEqual(data["error"]["step"], "docx")
        self.assertIn("requirements-vi.txt", data["error"]["fix"])

    def test_write_failure_reports_write_step(self):
        self.write_source()
        with mock.patch.object(de_thi, "load_docx_build") as loader:
            loader.return_value.FILENAMES = {"de": "de-en.docx", "song-ngu": "de-song-ngu.docx", "dap-an": "dap-an.docx"}
            loader.return_value.build.side_effect = OSError("file đang mở trong Word")
            code, data = self.run_cli()
        self.assertEqual(code, 1)
        self.assertEqual(data["error"]["step"], "write")
        self.assertIn("Word", data["error"]["fix"])

    def test_unexpected_failure_still_prints_one_json_line(self):
        self.write_source()
        with mock.patch.object(de_thi, "load_docx_build") as loader:
            loader.return_value.FILENAMES = {"de": "de-en.docx", "song-ngu": "de-song-ngu.docx", "dap-an": "dap-an.docx"}
            loader.return_value.build.side_effect = ValueError("lỗi ngoài dự kiến")
            code, data = self.run_cli()
        self.assertEqual(code, 1)
        self.assertEqual(data["error"]["step"], "internal")

    def test_missing_vietnamese_line_surfaces_as_a_warning(self):
        self.write_source(source_without("vi: Một vật khối lượng"))
        code, data = self.run_cli()
        self.assertEqual(code, 0, data)
        self.assertTrue(any("bản tiếng Việt" in warning for warning in data["warnings"]))

    def test_second_run_warns_about_overwriting(self):
        self.write_source()
        self.run_cli()
        code, data = self.run_cli()
        self.assertEqual(code, 0, data)
        self.assertTrue(any("Ghi đè" in warning for warning in data["warnings"]))

    def test_emit_falls_back_to_utf8_buffer(self):
        class LegacyStdout:
            def __init__(self):
                self.buffer = io.BytesIO()

            def write(self, text):
                text.encode("cp1252")
                return len(text)

            def flush(self):
                pass

        stream = LegacyStdout()
        with mock.patch.object(de_thi.sys, "stdout", stream):
            de_thi.emit({"ready": False, "warnings": ["Thiếu bản tiếng Việt"]})
        data = json.loads(stream.buffer.getvalue().decode("utf-8").strip())
        self.assertEqual(data["warnings"], ["Thiếu bản tiếng Việt"])


if __name__ == "__main__":
    unittest.main()
