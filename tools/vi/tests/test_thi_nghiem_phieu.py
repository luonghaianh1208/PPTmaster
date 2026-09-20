"""Test cho phiếu học tập Word của thí nghiệm ảo."""

import sys
import tempfile
import unittest
from pathlib import Path

TOOLS_VI = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(TOOLS_VI))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from thi_nghiem_mau import PENDULUM  # noqa: E402
from thi_nghiem_parts import parse, thu_vien  # noqa: E402


class WorksheetTest(unittest.TestCase):
    def setUp(self):
        from docx import Document
        from thi_nghiem_parts import phieu
        self.Document, self.phieu = Document, phieu
        self.model = thu_vien.load("li-con-lac-don", Path("."))
        self.experiment = parse.parse_experiment(PENDULUM, self.model.khai_bao)

    def test_ideal_table_uses_the_reference_model(self):
        rows, warning = self.phieu.ideal_table(self.experiment, self.model)
        self.assertEqual(warning, "")
        self.assertEqual([row[0] for row in rows], ["0,40", "0,80", "1,00", "1,40", "1,60"])
        self.assertEqual(rows[2][1], "2,007")

    def test_document_has_student_table_grid_and_teacher_page(self):
        with tempfile.TemporaryDirectory() as tmp:
            path, warnings = self.phieu.build(self.experiment, self.model, Path(tmp))
            document = self.Document(str(path))
        self.assertEqual(warnings, [])
        student, graph, ideal = document.tables
        self.assertEqual((len(student.rows), len(student.columns)), (6, 3))
        self.assertEqual([cell.text for cell in student.rows[0].cells], ["Lần", "Chiều dài dây l (m)", "Chu kì T (s)"])
        self.assertEqual((len(graph.rows), len(graph.columns)), (self.phieu.GRID_ROWS, self.phieu.GRID_COLUMNS))
        self.assertEqual(len(ideal.rows), 6)
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        for needle in ("PHIẾU HỌC TẬP", "DÀNH CHO GIÁO VIÊN", "Đáp án phần dự đoán: B", "Trục tung: (Chu kì T (s))2", "T = 2π√(l/g)"):
            self.assertIn(needle, text)
        self.assertLess(text.index("Dự đoán của em"), text.index("DÀNH CHO GIÁO VIÊN"))

    def test_no_slider_column_means_no_ideal_table(self):
        self.experiment.cot = ["chu-ki"]
        rows, warning = self.phieu.ideal_table(self.experiment, self.model)
        self.assertIsNone(rows)
        self.assertIn("thanh trượt", warning)

if __name__ == "__main__":
    unittest.main()
