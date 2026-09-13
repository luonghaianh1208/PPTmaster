"""Test cho lớp soạn giáo án tích hợp năng lực số và năng lực AI của bản Việt."""

import contextlib
import io
import json
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest import mock

REPO_ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(REPO_ROOT / "tools" / "vi"))

from giao_an_parts import docx_build, frameworks, parse, sgk  # noqa: E402

import giao_an  # noqa: E402

NLS_EXPECTED = (
    "1.1.NC1a", "1.1.NC1b", "1.2.NC1a", "1.3.NC1a",
    "2.1.NC1a", "2.2.NC1a",
    "3.1.NC1a", "3.2.NC1a",
    "4.1.NC1a",
    "5.1.NC1a", "5.3.NC1a",
)


class FrameworkLoadTest(unittest.TestCase):
    def setUp(self):
        self.frameworks = frameworks.load_file()

    def test_document_is_the_only_source_of_codes(self):
        self.assertTrue(frameworks.DOC_PATH.is_file(), frameworks.DOC_PATH)
        text = frameworks.DOC_PATH.read_text(encoding="utf-8")
        for code in NLS_EXPECTED:
            self.assertIn(f"`{code}`", text)

    def test_every_digital_competence_code_is_read(self):
        self.assertEqual(tuple(self.frameworks.nls), NLS_EXPECTED)

    def test_four_chemistry_ai_frameworks_are_read(self):
        self.assertEqual(
            sorted(self.frameworks.ai),
            sorted(["AI-H10", "AI-H11", "AI-H12", "AI-H-Chuyên"]),
        )

    def test_grade_eleven_framework_has_four_codes(self):
        self.assertEqual(
            self.frameworks.ai["AI-H11"],
            ["AI-H11.1", "AI-H11.2", "AI-H11.3", "AI-H11.4"],
        )

    def test_framework_is_chosen_by_subject_and_grade(self):
        self.assertEqual(self.frameworks.framework_for("Hoá học", "11"), "AI-H11")
        self.assertEqual(self.frameworks.framework_for("Hoá học", "10"), "AI-H10")
        self.assertEqual(self.frameworks.framework_for("Hoá học", "12"), "AI-H12")

    def test_specialised_track_uses_its_own_framework(self):
        self.assertEqual(self.frameworks.framework_for("Hoá học", "11", "chuyen"), "AI-H-Chuyên")
        self.assertEqual(self.frameworks.framework_for("Hoá học", "10", "hệ chuyên"), "AI-H-Chuyên")

    def test_subject_name_matching_ignores_case_and_diacritics(self):
        self.assertEqual(self.frameworks.framework_for("HOA HOC", "11"), "AI-H11")
        self.assertEqual(self.frameworks.framework_for("  hoá   học ", "11"), "AI-H11")

    def test_subject_without_a_framework_returns_none(self):
        self.assertIsNone(self.frameworks.framework_for("Ngữ văn", "11"))
        self.assertEqual(self.frameworks.ai_codes_for("Ngữ văn", "11"), [])

    def test_unknown_grade_returns_none(self):
        self.assertIsNone(self.frameworks.framework_for("Hoá học", "9"))

    def test_load_rejects_a_document_without_codes(self):
        with self.assertRaises(frameworks.FrameworkError):
            frameworks.load("# Trống\n\n## Khung năng lực số (mọi môn)\n")

    def test_normalise_strips_diacritics_and_case(self):
        self.assertEqual(frameworks.normalise("Hoá  Học"), "hoa hoc")
        self.assertEqual(frameworks.normalise("Hệ Chuyên"), "he chuyen")


VALID_SOURCE = """---
school: TRƯỜNG THPT CHUYÊN NGUYỄN TRÃI
department: TỔ HOÁ HỌC
teacher: Lương Hải Anh
subject: Hoá học
grade: 11
lesson: Bài 5. Ammonia và một số hợp chất ammonium
periods: 2
week: 5-6
period_numbers: 9-10
school_year: 2026-2027
---

## MUC TIEU

### Kien thuc
- Trình bày được tính chất hoá học của NH~3~ và muối ammonium.

### Nang luc chung
- Tự chủ và tự học: đọc SGK và tài liệu số để rút ra kết luận.

### Nang luc dac thu
- Nhận thức hoá học: giải thích tính base yếu của NH~3~.

### Nang luc so
- 1.1.NC1a — Tìm kiếm, chọn lọc dữ liệu về ứng dụng của ammonia trong sản xuất phân bón.
- 5.1.NC1a — Dùng thí nghiệm ảo PhET kiểm chứng chuyển dịch cân bằng.

### Nang luc AI
- AI-H11.2 — Viết prompt để AI dự đoán chiều chuyển dịch cân bằng khi tăng áp suất.
- AI-H11.4 — Đối chiếu kết quả AI với thí nghiệm, chỉ ra chỗ AI nói sai.

### Pham chat
- Trung thực: ghi đúng số liệu thí nghiệm, kể cả khi lệch dự đoán.

## THIET BI

### Giao vien
- Bộ dụng cụ điều chế NH~3~, máy chiếu, phiếu học tập 1.

### Hoc sinh
- SGK, vở, điện thoại có mạng theo nhóm.

## TIEN TRINH

### Mở đầu
thoi-luong: 10
muc-tieu: Tạo tình huống có vấn đề về mùi khai của ammonia.
noi-dung: Học sinh quan sát video và nêu dự đoán.
san-pham: Câu trả lời của học sinh.
chuyen-giao: Giáo viên chiếu video và nêu câu hỏi.
thuc-hien: Học sinh thảo luận cặp đôi trong 3 phút.
bao-cao: Hai nhóm trình bày dự đoán.
ket-luan: Giáo viên chốt vấn đề của bài học.

### Hình thành kiến thức — cân bằng tổng hợp ammonia
thoi-luong: 35
muc-tieu: Dự đoán và kiểm chứng chiều chuyển dịch cân bằng.
noi-dung: Học sinh viết prompt cho AI dự đoán chiều chuyển dịch khi tăng áp suất.
noi-dung: Sau đó dùng PhET để kiểm chứng và hoàn thành Phiếu học tập 1.
san-pham: Đoạn chat với AI và bảng số liệu trong phiếu học tập 1.
nls: 1.1.NC1a, 5.1.NC1a
ai: AI-H11.2, AI-H11.4
chuyen-giao: Giáo viên giao mẫu prompt và địa chỉ mô phỏng PhET.
thuc-hien: Học sinh làm theo nhóm 4, giáo viên quan sát.
bao-cao: Mỗi nhóm nêu một chỗ AI trả lời chưa đúng.
ket-luan: Giáo viên chuẩn hoá nguyên lý Le Chatelier.

### Luyện tập
thoi-luong: 30
muc-tieu: Vận dụng tính chất của muối ammonium.
noi-dung: Học sinh làm bài tập trong phiếu học tập 1.
san-pham: Bài giải trong phiếu.
chuyen-giao: Giáo viên phát phiếu.
thuc-hien: Học sinh làm cá nhân.
bao-cao: Ba học sinh lên bảng.
ket-luan: Giáo viên nhận xét và chữa bài.

### Vận dụng
thoi-luong: 15
muc-tieu: Liên hệ sản xuất phân bón ở địa phương.
noi-dung: Học sinh đề xuất cách bảo quản phân đạm.
san-pham: Đoạn trả lời ngắn.
chuyen-giao: Giáo viên giao nhiệm vụ về nhà.
thuc-hien: Học sinh làm ở nhà.
bao-cao: Nộp qua lớp học trực tuyến.
ket-luan: Giáo viên nhận xét ở tiết sau.

## PHIEU HOC TAP

### Phiếu học tập 1
- Câu 1. Viết phương trình tổng hợp NH~3~ từ N~2~ và H~2~.
- Câu 2. Dự đoán chiều chuyển dịch cân bằng khi tăng áp suất.

## RUBRIC

### Kĩ năng viết prompt
- Mức 1: Viết được câu lệnh nhưng thiếu dữ kiện.
- Mức 2: Viết được câu lệnh đủ dữ kiện.
- Mức 3: Viết được câu lệnh đủ dữ kiện và có ràng buộc rõ ràng.

### Kĩ năng kiểm chứng và phản biện AI
- Mức 1: Chấp nhận kết quả AI mà không kiểm chứng.
- Mức 2: Kiểm chứng được bằng thí nghiệm ảo.
- Mức 3: Kiểm chứng và chỉ ra được chỗ AI nói sai kèm lý do.

## CAN SOAT
- Thầy cô xác nhận giúp em địa chỉ mô phỏng PhET dùng được ở phòng máy của trường.
"""


def source_without(prefix: str) -> str:
    kept = [line for line in VALID_SOURCE.splitlines() if not line.startswith(prefix)]
    return "\n".join(kept) + "\n"


def line_of(text: str, prefix: str, occurrence: int = 1) -> int:
    """Số dòng (đếm từ 1) của lần thứ `occurrence` một dòng mở đầu bằng `prefix`."""
    seen = 0
    for number, line in enumerate(text.splitlines(), start=1):
        if line.startswith(prefix):
            seen += 1
            if seen == occurrence:
                return number
    raise AssertionError(f"không tìm thấy dòng {prefix!r}")


class ParseTest(unittest.TestCase):
    def test_valid_source_reads_every_section(self):
        lesson = parse.parse_lesson(VALID_SOURCE)
        self.assertEqual(lesson.meta["subject"], "Hoá học")
        self.assertEqual(lesson.meta["grade"], "11")
        self.assertEqual(len(lesson.activities), 4)
        self.assertEqual(len(lesson.worksheets), 1)
        self.assertEqual(len(lesson.rubrics), 2)
        self.assertEqual(len(lesson.review_notes), 1)
        self.assertFalse(lesson.ai_no_framework)

    def test_codes_are_read_from_the_objective_lines(self):
        lesson = parse.parse_lesson(VALID_SOURCE)
        self.assertEqual(lesson.nls_codes, ["1.1.NC1a", "5.1.NC1a"])
        self.assertEqual(lesson.ai_codes, ["AI-H11.2", "AI-H11.4"])

    def test_repeated_key_becomes_two_paragraphs(self):
        activity = parse.parse_lesson(VALID_SOURCE).activities[1]
        self.assertEqual(len(activity.noi_dung), 2)
        self.assertIn("PhET", activity.noi_dung[1])

    def test_activity_codes_are_split_on_commas(self):
        activity = parse.parse_lesson(VALID_SOURCE).activities[1]
        self.assertEqual(activity.nls, ["1.1.NC1a", "5.1.NC1a"])
        self.assertEqual(activity.ai, ["AI-H11.2", "AI-H11.4"])

    def test_rubric_levels_are_stripped_of_their_prefix(self):
        rubric = parse.parse_lesson(VALID_SOURCE).rubrics[0]
        self.assertEqual(len(rubric.levels), 3)
        self.assertTrue(rubric.levels[0].startswith("Viết được câu lệnh nhưng"))

    def test_missing_meta_key_names_it(self):
        with self.assertRaises(parse.ParseError) as caught:
            parse.parse_lesson(source_without("subject:"))
        self.assertIn("subject", caught.exception.message)

    def test_grade_must_be_digits(self):
        broken = VALID_SOURCE.replace("grade: 11", "grade: mười một")
        with self.assertRaises(parse.ParseError) as caught:
            parse.parse_lesson(broken)
        self.assertIn("grade", caught.exception.message)
        self.assertEqual(caught.exception.line_no, line_of(broken, "grade:"))

    def test_periods_must_be_a_positive_number(self):
        broken = VALID_SOURCE.replace("periods: 2", "periods: 0")
        with self.assertRaises(parse.ParseError) as caught:
            parse.parse_lesson(broken)
        self.assertIn("periods", caught.exception.message)
        self.assertEqual(caught.exception.line_no, line_of(broken, "periods:"))

    def test_activity_missing_any_of_the_eight_keys_fails(self):
        """Bốn thành tố 5512 và bốn bước tổ chức là thứ người kiểm tra soi trước tiên."""
        for key in ("thoi-luong", "muc-tieu", "noi-dung", "san-pham",
                    "chuyen-giao", "thuc-hien", "bao-cao", "ket-luan"):
            with self.subTest(key=key):
                dropped = False
                kept = []
                for line in VALID_SOURCE.splitlines():
                    if not dropped and line.startswith(f"{key}:"):
                        dropped = True
                        continue
                    kept.append(line)
                self.assertTrue(dropped, f"không tìm thấy dòng {key}:")
                with self.assertRaises(parse.ParseError) as caught:
                    parse.parse_lesson("\n".join(kept) + "\n")
                self.assertIn(key, caught.exception.message)

    def test_activity_missing_muc_tieu_names_the_key_and_line(self):
        broken = VALID_SOURCE.replace("muc-tieu: Tạo tình huống có vấn đề về mùi khai của ammonia.\n", "", 1)
        with self.assertRaises(parse.ParseError) as caught:
            parse.parse_lesson(broken)
        self.assertIn("muc-tieu", caught.exception.message)
        self.assertEqual(caught.exception.line_no, line_of(broken, "### Mở đầu"))

    def test_activity_time_must_be_positive(self):
        broken = VALID_SOURCE.replace("thoi-luong: 10", "thoi-luong: 0", 1)
        with self.assertRaises(parse.ParseError) as caught:
            parse.parse_lesson(broken)
        self.assertIn("thoi-luong", caught.exception.message)
        self.assertEqual(caught.exception.line_no, line_of(broken, "thoi-luong: 0"))

    def test_unknown_activity_key_fails(self):
        broken = VALID_SOURCE.replace("bao-cao: Hai nhóm trình bày dự đoán.", "ghi-chu: abc", 1)
        with self.assertRaises(parse.ParseError) as caught:
            parse.parse_lesson(broken)
        self.assertIn("ghi-chu", caught.exception.message)

    def test_sections_out_of_order_fail(self):
        broken = VALID_SOURCE.replace("## THIET BI", "## RUBRIC", 1)
        with self.assertRaises(parse.ParseError) as caught:
            parse.parse_lesson(broken)
        self.assertIn("thứ tự", caught.exception.message)
        self.assertEqual(caught.exception.line_no, line_of(broken, "## TIEN TRINH"))

    def test_duplicate_section_fails(self):
        broken = VALID_SOURCE + "\n## RUBRIC\n\n### Thêm\n- Mức 1: a\n- Mức 2: b\n- Mức 3: c\n"
        with self.assertRaises(parse.ParseError) as caught:
            parse.parse_lesson(broken)
        self.assertIn("hai lần", caught.exception.message)
        self.assertEqual(caught.exception.line_no, line_of(broken, "## RUBRIC", 2))

    def test_review_section_must_be_last(self):
        broken = VALID_SOURCE.replace("## RUBRIC", "## CAN SOAT\n- ghi chú\n\n## RUBRIC", 1)
        with self.assertRaises(parse.ParseError) as caught:
            parse.parse_lesson(broken)
        self.assertIn("thứ tự", caught.exception.message)
        self.assertEqual(caught.exception.line_no, line_of(broken, "## RUBRIC"))

    def test_objective_needs_all_six_subsections_in_order(self):
        broken = VALID_SOURCE.replace("### Pham chat\n- Trung thực: ghi đúng số liệu thí nghiệm, kể cả khi lệch dự đoán.\n", "")
        with self.assertRaises(parse.ParseError) as caught:
            parse.parse_lesson(broken)
        self.assertIn("Pham chat", caught.exception.message)

    def test_empty_digital_competence_section_fails(self):
        broken = source_without("- 1.1.NC1a").replace("- 5.1.NC1a — Dùng thí nghiệm ảo PhET kiểm chứng chuyển dịch cân bằng.\n", "")
        with self.assertRaises(parse.ParseError) as caught:
            parse.parse_lesson(broken)
        self.assertIn("Nang luc so", caught.exception.message)

    def test_empty_ai_competence_section_fails(self):
        broken = VALID_SOURCE.replace(
            "- AI-H11.2 — Viết prompt để AI dự đoán chiều chuyển dịch cân bằng khi tăng áp suất.\n"
            "- AI-H11.4 — Đối chiếu kết quả AI với thí nghiệm, chỉ ra chỗ AI nói sai.\n",
            ""
        )
        with self.assertRaises(parse.ParseError) as caught:
            parse.parse_lesson(broken)
        self.assertIn("Nang luc AI", caught.exception.message)

    def test_no_activity_carrying_a_digital_code_fails(self):
        broken = VALID_SOURCE.replace("nls: 1.1.NC1a, 5.1.NC1a\n", "")
        with self.assertRaises(parse.ParseError) as caught:
            parse.parse_lesson(broken)
        self.assertIn("nls", caught.exception.message)
        self.assertIn("tiến trình", caught.exception.message)

    def test_no_activity_carrying_an_ai_code_fails(self):
        broken = VALID_SOURCE.replace("ai: AI-H11.2, AI-H11.4\n", "")
        with self.assertRaises(parse.ParseError) as caught:
            parse.parse_lesson(broken)
        self.assertIn("ai", caught.exception.message)
        self.assertIn("tiến trình", caught.exception.message)

    def test_exam_with_no_activity_fails(self):
        # Giữ dòng '## PHIEU HOC TAP' để lỗi đúng là "không có hoạt động", không phải phiếu bị đọc thành hoạt động.
        broken = (VALID_SOURCE.split("## TIEN TRINH")[0] + "## TIEN TRINH\n\n## PHIEU HOC TAP"
                  + VALID_SOURCE.split("## PHIEU HOC TAP")[1])
        with self.assertRaises(parse.ParseError) as caught:
            parse.parse_lesson(broken)
        self.assertIn("không có hoạt động nào", caught.exception.message)
        self.assertEqual(caught.exception.line_no, line_of(broken, "## TIEN TRINH"))

    def test_rubric_needs_at_least_two_criteria(self):
        broken = VALID_SOURCE.split("### Kĩ năng kiểm chứng và phản biện AI")[0] + "\n## CAN SOAT\n- ghi chú\n"
        with self.assertRaises(parse.ParseError) as caught:
            parse.parse_lesson(broken)
        self.assertIn("hai tiêu chí", caught.exception.message)

    def test_rubric_criterion_needs_exactly_three_levels(self):
        broken = VALID_SOURCE.replace("- Mức 3: Viết được câu lệnh đủ dữ kiện và có ràng buộc rõ ràng.\n", "")
        with self.assertRaises(parse.ParseError) as caught:
            parse.parse_lesson(broken)
        self.assertIn("ba mức", caught.exception.message)

    def test_rubric_levels_must_be_in_order(self):
        broken = VALID_SOURCE.replace("- Mức 1: Viết được câu lệnh nhưng thiếu dữ kiện.", "- Mức 2: Viết được câu lệnh nhưng thiếu dữ kiện.", 1)
        with self.assertRaises(parse.ParseError) as caught:
            parse.parse_lesson(broken)
        self.assertIn("Mức 1:", caught.exception.message)

    def test_subject_without_a_framework_declares_the_marker(self):
        broken = VALID_SOURCE.replace("subject: Hoá học", "subject: Ngữ văn")
        broken = broken.replace(
            "- AI-H11.2 — Viết prompt để AI dự đoán chiều chuyển dịch cân bằng khi tăng áp suất.\n"
            "- AI-H11.4 — Đối chiếu kết quả AI với thí nghiệm, chỉ ra chỗ AI nói sai.\n",
            f"- {frameworks.NO_FRAMEWORK}\n- Học sinh nhận xét đoạn văn do AI viết.\n",
        )
        broken = broken.replace("ai: AI-H11.2, AI-H11.4", f"ai: {frameworks.NO_CODE}")
        lesson = parse.parse_lesson(broken)
        self.assertTrue(lesson.ai_no_framework)
        self.assertEqual(lesson.ai_codes, [])
        self.assertEqual(lesson.activities[1].ai, [frameworks.NO_CODE])

    def test_digital_competence_line_must_use_em_dash(self):
        broken = VALID_SOURCE.replace(
            "- 1.1.NC1a — Tìm kiếm, chọn lọc dữ liệu về ứng dụng của ammonia trong sản xuất phân bón.",
            "- 1.1.NC1a – Tìm kiếm, chọn lọc dữ liệu về ứng dụng của ammonia trong sản xuất phân bón.",
        )
        with self.assertRaises(parse.ParseError) as caught:
            parse.parse_lesson(broken)
        self.assertIn("—", caught.exception.message)
        self.assertEqual(caught.exception.line_no, line_of(broken, "- 1.1.NC1a –"))

    def test_digital_competence_line_without_a_code_fails(self):
        broken = VALID_SOURCE.replace(
            "- 1.1.NC1a — Tìm kiếm, chọn lọc dữ liệu về ứng dụng của ammonia trong sản xuất phân bón.",
            "- Tìm kiếm thông tin",
        )
        with self.assertRaises(parse.ParseError) as caught:
            parse.parse_lesson(broken)
        self.assertEqual(caught.exception.line_no, line_of(broken, "- Tìm kiếm thông tin"))

    def test_track_must_be_a_recognised_specialised_alias(self):
        broken = VALID_SOURCE.replace("grade: 11", "grade: 11\ntrack: thuong")
        with self.assertRaises(parse.ParseError) as caught:
            parse.parse_lesson(broken)
        self.assertIn("track", caught.exception.message)
        self.assertEqual(caught.exception.line_no, line_of(broken, "track:"))


class LessonMathTest(unittest.TestCase):
    def test_total_time_matches_two_periods(self):
        lesson = parse.parse_lesson(VALID_SOURCE)
        self.assertEqual(lesson.minutes, 90)
        self.assertEqual(lesson.periods, 2)
        self.assertEqual([warning for warning in lesson.warnings if "thời lượng" in warning], [])

    def test_time_mismatch_is_a_warning(self):
        broken = VALID_SOURCE.replace("thoi-luong: 30", "thoi-luong: 20", 1)
        lesson = parse.parse_lesson(broken)
        self.assertTrue(any("thời lượng" in warning for warning in lesson.warnings))

    def test_missing_worksheet_section_is_a_warning(self):
        broken = VALID_SOURCE.split("## PHIEU HOC TAP")[0] + "## RUBRIC" + VALID_SOURCE.split("## RUBRIC")[1]
        lesson = parse.parse_lesson(broken)
        self.assertTrue(any("PHIEU HOC TAP" in warning for warning in lesson.warnings))

    def test_worksheet_mentioned_but_not_defined_is_a_warning(self):
        broken = VALID_SOURCE.replace("### Phiếu học tập 1", "### Phiếu học tập 2")
        lesson = parse.parse_lesson(broken)
        self.assertTrue(any("Phiếu học tập 1" in warning for warning in lesson.warnings))

    def test_worksheet_mention_with_so_word_is_recognised(self):
        broken = VALID_SOURCE.replace(
            "san-pham: Đoạn chat với AI và bảng số liệu trong phiếu học tập 1.",
            "san-pham: Đoạn chat với AI và bảng số liệu trong Phiếu học tập số 2.",
        )
        lesson = parse.parse_lesson(broken)
        self.assertTrue(any("Phiếu học tập 2" in warning for warning in lesson.warnings))


class FrameworkValidateTest(unittest.TestCase):
    def setUp(self):
        self.frameworks = frameworks.load_file()

    def lesson_from(self, text: str):
        return parse.parse_lesson(text)

    def test_valid_lesson_passes(self):
        frameworks.validate(self.lesson_from(VALID_SOURCE), self.frameworks)

    def test_unknown_digital_code_is_rejected(self):
        broken = VALID_SOURCE.replace("1.1.NC1a —", "9.9.NC9z —")
        broken = broken.replace("nls: 1.1.NC1a,", "nls: 9.9.NC9z,")
        with self.assertRaises(frameworks.FrameworkError) as caught:
            frameworks.validate(self.lesson_from(broken), self.frameworks)
        self.assertIn("9.9.NC9z", caught.exception.message)
        self.assertIn("1.1.NC1a", caught.exception.fix)

    def test_ai_code_outside_the_grade_framework_is_rejected(self):
        broken = VALID_SOURCE.replace("AI-H11.2 —", "AI-H12 —")
        broken = broken.replace("ai: AI-H11.2,", "ai: AI-H12,")
        with self.assertRaises(frameworks.FrameworkError) as caught:
            frameworks.validate(self.lesson_from(broken), self.frameworks)
        self.assertIn("AI-H12", caught.exception.message)
        self.assertIn("AI-H11", caught.exception.message)

    def test_subject_without_a_framework_must_use_the_marker(self):
        broken = VALID_SOURCE.replace("subject: Hoá học", "subject: Ngữ văn")
        with self.assertRaises(frameworks.FrameworkError) as caught:
            frameworks.validate(self.lesson_from(broken), self.frameworks)
        self.assertIn(frameworks.NO_FRAMEWORK, caught.exception.fix)

    def test_subject_with_a_framework_must_not_use_the_marker(self):
        broken = VALID_SOURCE.replace(
            "- AI-H11.2 — Viết prompt để AI dự đoán chiều chuyển dịch cân bằng khi tăng áp suất.\n"
            "- AI-H11.4 — Đối chiếu kết quả AI với thí nghiệm, chỉ ra chỗ AI nói sai.\n",
            f"- {frameworks.NO_FRAMEWORK}\n- Học sinh đối chiếu kết quả AI.\n",
        ).replace("ai: AI-H11.2, AI-H11.4", f"ai: {frameworks.NO_CODE}")
        with self.assertRaises(frameworks.FrameworkError) as caught:
            frameworks.validate(self.lesson_from(broken), self.frameworks)
        self.assertIn("AI-H11", caught.exception.message)

    def test_subject_with_a_framework_must_not_use_the_no_code_marker_in_activities(self):
        broken = VALID_SOURCE.replace("ai: AI-H11.2, AI-H11.4", f"ai: {frameworks.NO_CODE}")
        with self.assertRaises(frameworks.FrameworkError) as caught:
            frameworks.validate(self.lesson_from(broken), self.frameworks)
        self.assertIn(frameworks.NO_CODE, caught.exception.message)
        self.assertIn("AI-H11.2", caught.exception.fix)

    def test_marker_subject_passes_with_no_code_activities(self):
        text = VALID_SOURCE.replace("subject: Hoá học", "subject: Ngữ văn")
        text = text.replace(
            "- AI-H11.2 — Viết prompt để AI dự đoán chiều chuyển dịch cân bằng khi tăng áp suất.\n"
            "- AI-H11.4 — Đối chiếu kết quả AI với thí nghiệm, chỉ ra chỗ AI nói sai.\n",
            f"- {frameworks.NO_FRAMEWORK}\n- Học sinh nhận xét đoạn văn do AI viết.\n",
        )
        text = text.replace("ai: AI-H11.2, AI-H11.4", f"ai: {frameworks.NO_CODE}")
        frameworks.validate(self.lesson_from(text), self.frameworks)

    def test_marker_subject_may_not_carry_codes(self):
        text = VALID_SOURCE.replace("subject: Hoá học", "subject: Ngữ văn")
        text = text.replace(
            "- AI-H11.2 — Viết prompt để AI dự đoán chiều chuyển dịch cân bằng khi tăng áp suất.\n",
            f"- {frameworks.NO_FRAMEWORK}\n",
        )
        with self.assertRaises(frameworks.FrameworkError) as caught:
            frameworks.validate(self.lesson_from(text), self.frameworks)
        self.assertIn("AI-H11.4", caught.exception.message)

    def test_marker_subject_objectives_may_not_carry_codes(self):
        """Spec §7: môn chưa có khung mà mục tiêu AI vẫn ghi mã thì chặn, kể cả khi hoạt động đã dùng (chưa có mã)."""
        text = VALID_SOURCE.replace("subject: Hoá học", "subject: Ngữ văn")
        text = text.replace(
            "- AI-H11.2 — Viết prompt để AI dự đoán chiều chuyển dịch cân bằng khi tăng áp suất.\n",
            f"- {frameworks.NO_FRAMEWORK}\n",
        )
        text = text.replace("ai: AI-H11.2, AI-H11.4", f"ai: {frameworks.NO_CODE}")
        with self.assertRaises(frameworks.FrameworkError) as caught:
            frameworks.validate(self.lesson_from(text), self.frameworks)
        self.assertIn("AI-H11.4", caught.exception.message)
        self.assertIn(frameworks.NO_FRAMEWORK, caught.exception.fix)

    def test_activity_code_missing_from_the_objectives_is_rejected(self):
        broken = VALID_SOURCE.replace("nls: 1.1.NC1a, 5.1.NC1a", "nls: 1.1.NC1a, 3.1.NC1a")
        with self.assertRaises(frameworks.FrameworkError) as caught:
            frameworks.validate(self.lesson_from(broken), self.frameworks)
        self.assertIn("3.1.NC1a", caught.exception.message)
        self.assertIn("Nang luc so", caught.exception.message)

    def test_activity_ai_code_missing_from_the_objectives_is_rejected(self):
        broken = VALID_SOURCE.replace("ai: AI-H11.2, AI-H11.4", "ai: AI-H11.2, AI-H11.3")
        with self.assertRaises(frameworks.FrameworkError) as caught:
            frameworks.validate(self.lesson_from(broken), self.frameworks)
        self.assertIn("AI-H11.3", caught.exception.message)
        self.assertIn("Nang luc AI", caught.exception.message)

    def test_specialised_track_uses_the_specialised_codes(self):
        text = VALID_SOURCE.replace("grade: 11", "grade: 11\ntrack: chuyen")
        text = text.replace("- AI-H11.2 —", "- AI-H-Chuyên —")
        text = text.replace("- AI-H11.4 — Đối chiếu kết quả AI với thí nghiệm, chỉ ra chỗ AI nói sai.\n", "")
        text = text.replace("ai: AI-H11.2, AI-H11.4", "ai: AI-H-Chuyên")
        frameworks.validate(self.lesson_from(text), self.frameworks)

    def test_subject_typo_lists_the_available_frameworks(self):
        broken = VALID_SOURCE.replace("subject: Hoá học", "subject: Hoá")
        with self.assertRaises(frameworks.FrameworkError) as caught:
            frameworks.validate(self.lesson_from(broken), self.frameworks)
        self.assertIn("Hoá học", caught.exception.message)
        self.assertIn("lớp 11", caught.exception.message)
        self.assertIn("subject:", caught.exception.fix)

    def test_activity_nls_code_unknown_to_the_ministry_is_rejected(self):
        broken = VALID_SOURCE.replace("nls: 1.1.NC1a, 5.1.NC1a", "nls: 1.1.NC1a, 5.1.NC1b")
        with self.assertRaises(frameworks.FrameworkError) as caught:
            frameworks.validate(self.lesson_from(broken), self.frameworks)
        self.assertIn("5.1.NC1b", caught.exception.message)
        self.assertIn("1.1.NC1a", caught.exception.fix)

    def test_activity_ai_code_outside_the_activitys_own_framework_is_rejected(self):
        broken = VALID_SOURCE.replace("ai: AI-H11.2, AI-H11.4", "ai: AI-H11.2, AI-H12")
        with self.assertRaises(frameworks.FrameworkError) as caught:
            frameworks.validate(self.lesson_from(broken), self.frameworks)
        self.assertIn("AI-H12", caught.exception.message)
        self.assertIn("AI-H11.4", caught.exception.fix)


SGK_SAMPLE = """# Chương 2. Nitrogen và sulfur

Mở đầu chương.

## Bài 4. Nitrogen

Nội dung bài 4.

## Bài 5. Ammonia và một số hợp chất ammonium

Nội dung bài 5, đoạn một.

Nội dung bài 5, đoạn hai.

## Bài 6. Một số hợp chất với oxygen của nitrogen

Nội dung bài 6.
"""


class SgkTest(unittest.TestCase):
    def test_headings_are_found(self):
        titles = [title for _, title in sgk.headings(SGK_SAMPLE.splitlines())]
        self.assertIn("Bài 5. Ammonia và một số hợp chất ammonium", titles)
        self.assertEqual(len(titles), 4)

    def test_extract_takes_only_the_requested_lesson(self):
        heading, body, warnings = sgk.extract(SGK_SAMPLE, "Bài 5")
        self.assertTrue(heading.startswith("Bài 5."))
        self.assertIn("đoạn một", body)
        self.assertIn("đoạn hai", body)
        self.assertNotIn("Nội dung bài 6", body)
        self.assertNotIn("Nội dung bài 4", body)
        self.assertEqual(warnings, [])

    def test_extract_ignores_case_and_diacritics(self):
        heading, _, _ = sgk.extract(SGK_SAMPLE, "bai 5")
        self.assertTrue(heading.startswith("Bài 5."))

    def test_unknown_lesson_names_the_nearest_headings(self):
        with self.assertRaises(sgk.SgkError) as caught:
            sgk.extract(SGK_SAMPLE, "Bài 99")
        self.assertIn("Bài 99", caught.exception.message)
        self.assertIn("gần nhất", caught.exception.fix)
        self.assertIn("Bài 4.", caught.exception.fix)

    def test_file_without_headings_is_rejected(self):
        with self.assertRaises(sgk.SgkError) as caught:
            sgk.extract("chỉ là văn bản thường\nkhông có tiêu đề\n", "Bài 5")
        self.assertIn("tiêu đề", caught.exception.message)

    def test_plain_bai_lines_count_as_headings(self):
        text = "BÀI 5. AMMONIA\n\nnội dung\n\nBÀI 6. NITRIC ACID\n\nkhác\n"
        heading, body, _ = sgk.extract(text, "Bài 5")
        self.assertIn("AMMONIA", heading)
        self.assertNotIn("NITRIC", body)

    def test_oversized_slice_warns(self):
        filler = "x" * 1024
        text = "## Bài 5. Ammonia\n\n" + "\n".join([filler] * 220) + "\n"
        _, _, warnings = sgk.extract(text, "Bài 5")
        self.assertTrue(any("KB" in warning for warning in warnings))

    def test_several_matches_warn_and_take_the_first(self):
        text = "## Bài 5. Ammonia\n\nmột\n\n## Bài 5. Ammonia (tiếp)\n\nhai\n"
        heading, body, warnings = sgk.extract(text, "Bài 5")
        self.assertEqual(heading, "Bài 5. Ammonia")
        self.assertIn("một", body)
        self.assertTrue(any("khớp" in warning for warning in warnings))

    def test_table_of_contents_lines_are_skipped(self):
        text = (
            "MỤC LỤC\nBài 4. Nitrogen 20\nBài 5. Ammonia 25\nBài 6. Nitric acid 30\n\n"
            "## Bài 5. Ammonia\n\nnội dung thật\n\n## Bài 6. Nitric acid\n\nkhác\n"
        )
        heading, body, _ = sgk.extract(text, "Bài 5")
        self.assertEqual(heading, "Bài 5. Ammonia")
        self.assertIn("nội dung thật", body)
        self.assertNotIn("khác", body)

    def test_lesson_number_does_not_match_a_longer_number(self):
        text = "## Bài 50. Hệ sinh thái\n\nsai\n\n## Bài 5. Tế bào\n\nđúng\n"
        heading, body, warnings = sgk.extract(text, "Bài 5")
        self.assertEqual(heading, "Bài 5. Tế bào")
        self.assertIn("đúng", body)
        self.assertEqual(warnings, [])

    def test_sub_headings_inside_a_lesson_are_kept(self):
        text = (
            "# Chương 2. Nitrogen và sulfur\n\n## Bài 5. Ammonia\n\nMở đầu bài.\n\n"
            "### I. Tính chất vật lí\n\nkhí không màu\n\n## II. Tính chất hoá học\n\ntính base yếu\n\n"
            "## Bài 6. Nitric acid\n\nkhác\n"
        )
        heading, body, warnings = sgk.extract(text, "Bài 5")
        self.assertEqual(heading, "Bài 5. Ammonia")
        self.assertIn("khí không màu", body)
        self.assertIn("tính base yếu", body)
        self.assertNotIn("khác", body)
        self.assertEqual(warnings, [])

    def test_a_shallower_heading_ends_the_lesson(self):
        text = "## Bài 8. Sulfuric acid\n\nnội dung bài 8\n\n# Chương 3. Đại cương hoá hữu cơ\n\nmở đầu chương 3\n"
        _, body, _ = sgk.extract(text, "Bài 8")
        self.assertIn("nội dung bài 8", body)
        self.assertNotIn("mở đầu chương 3", body)

    def test_extract_ignores_punctuation_in_the_query_and_the_heading(self):
        text = "## BÀI 5: AMMONIA\n\nnội dung thật\n\n## Bài 6. Nitric acid\n\nkhác\n"
        heading, body, _ = sgk.extract(text, "Bài 5. Ammonia")
        self.assertEqual(heading, "BÀI 5: AMMONIA")
        self.assertIn("nội dung thật", body)

    def test_table_of_contents_stray_chapter_line_is_not_mistaken_for_content(self):
        # (a) Mục lục có dòng chương thường ("CHƯƠNG 3...", không '#') xen giữa hai dòng
        # mục lục "Bài <số>" khiến lát cắt mục lục trông như "có nội dung".
        text = (
            "MỤC LỤC\nBài 4. Nitrogen 20\nBài 5. Ammonia 25\n"
            "CHƯƠNG 3. ĐẠI CƯƠNG HỮU CƠ\nBài 6. Nitric acid 30\n\n"
            "## Bài 5. Ammonia\n\nnội dung thật\n\n## Bài 6. Nitric acid\n\nkhác\n"
        )
        heading, body, _ = sgk.extract(text, "Bài 5")
        self.assertEqual(heading, "Bài 5. Ammonia")
        self.assertIn("nội dung thật", body)
        self.assertNotIn("CHƯƠNG 3", body)
        self.assertNotIn("khác", body)

    def test_exercise_numbering_inside_the_lesson_does_not_cut_it_short(self):
        # (b) "Bài 1. ..." là số thứ tự bài tập trong bài 5, không phải bài mới.
        text = (
            "## Bài 5. Ammonia\n\nNội dung lý thuyết.\n\n"
            "Bài 1. Viết phương trình\n\nBài 2. Tính\n\n"
            "## Bài 6. Nitric acid\n\nkhác\n"
        )
        heading, body, _ = sgk.extract(text, "Bài 5")
        self.assertEqual(heading, "Bài 5. Ammonia")
        self.assertIn("Bài 1. Viết phương trình", body)
        self.assertIn("Bài 2. Tính", body)
        self.assertNotIn("khác", body)

    def test_all_plain_heading_book_picks_the_longest_slice(self):
        text = (
            "Bài 4. Nitrogen\n"
            "  (trang 20)\n"
            "Bài 5. Ammonia\n"
            "  (trang 25)\n"
            "Bài 6. Nitric acid\n"
            "  (trang 30)\n"
            "Bài 5. Ammonia\n"
            "Nội dung thật rất dài của bài 5, mô tả chi tiết tính chất hoá học và "
            "ứng dụng thực tế trong nông nghiệp.\n"
            "Bài 6. Nitric acid\n"
            "nội dung khác\n"
        )
        heading, body, _ = sgk.extract(text, "Bài 5")
        self.assertIn("Nội dung thật", body)
        self.assertNotIn("trang 25", body)
        self.assertNotIn("nội dung khác", body)

    def test_last_lesson_of_a_plain_book_with_a_table_of_contents(self):
        text = (
            "MỤC LỤC\n"
            "Bài 5. Ammonia 25\n"
            "Bài 6. Nitric acid 32\n"
            "Bài 7. Hữu cơ 40\n"
            "\n"
            "BÀI 5. AMMONIA\n"
            "Nội dung lý thuyết bài 5.\n"
            "Tính chất hoá học của ammonia.\n"
            "ĐÁNH DẤU BÀI 5\n"
            "Ứng dụng trong nông nghiệp.\n"
            "\n"
            "BÀI 6. NITRIC ACID\n"
            "Nội dung lý thuyết bài 6.\n"
            "Tính chất hoá học của nitric acid.\n"
            "ĐÁNH DẤU BÀI 6\n"
            "Ứng dụng trong công nghiệp.\n"
            "\n"
            "BÀI 7. HỮU CƠ\n"
            "Nội dung lý thuyết bài 7.\n"
            "Tính chất hoá học của hợp chất hữu cơ.\n"
            "ĐÁNH DẤU BÀI 7\n"
            "Ứng dụng trong đời sống.\n"
        )
        heading, body, _ = sgk.extract(text, "Bài 7")
        self.assertEqual(heading, "BÀI 7. HỮU CƠ")
        self.assertIn("ĐÁNH DẤU BÀI 7", body)
        self.assertNotIn("ĐÁNH DẤU BÀI 5", body)
        self.assertNotIn("Bài 5. Ammonia 25", body)
        self.assertNotIn("Bài 6. Nitric acid 32", body)
        self.assertNotIn("Bài 7. Hữu cơ 40", body)

        heading5, body5, _ = sgk.extract(text, "Bài 5")
        self.assertEqual(heading5, "BÀI 5. AMMONIA")
        self.assertIn("ĐÁNH DẤU BÀI 5", body5)
        self.assertNotIn("ĐÁNH DẤU BÀI 6", body5)

        heading6, body6, _ = sgk.extract(text, "Bài 6")
        self.assertEqual(heading6, "BÀI 6. NITRIC ACID")
        self.assertIn("ĐÁNH DẤU BÀI 6", body6)
        self.assertNotIn("ĐÁNH DẤU BÀI 7", body6)

    def test_multi_match_warning_names_the_chosen_heading(self):
        text = (
            "MỤC LỤC\nBài 4. Nitrogen 20\nBài 5. Ammonia 25\nBài 6. Nitric acid 30\n\n"
            "## Bài 5. Ammonia\n\nnội dung thật\n\n## Bài 6. Nitric acid\n\nkhác\n"
        )
        heading, _, warnings = sgk.extract(text, "Bài 5")
        matched_warning = next(warning for warning in warnings if "khớp" in warning)
        self.assertIn(heading, matched_warning)
        self.assertNotIn("đầu tiên", matched_warning)


def document_xml(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        return archive.read("word/document.xml").decode("utf-8")


class DocxBuildTest(unittest.TestCase):
    def setUp(self):
        self.lesson = parse.parse_lesson(VALID_SOURCE)
        self.tmp = tempfile.TemporaryDirectory()
        self.folder = Path(self.tmp.name)
        self.addCleanup(self.tmp.cleanup)

    def test_page_uses_a4_and_the_school_margins(self):
        from docx import Document

        path = docx_build.build(self.lesson, self.folder)
        section = Document(str(path)).sections[0]
        # Word lưu khổ giấy theo twip nên đọc lại lệch vài trăm EMU (< 0,001 cm); so tới 0,01 cm.
        for attribute, expected_cm in (
            ("page_width", 21.0),
            ("page_height", 29.7),
            ("top_margin", 2.0),
            ("bottom_margin", 2.0),
            ("left_margin", 2.5),
            ("right_margin", 2.0),
        ):
            with self.subTest(attribute=attribute):
                self.assertAlmostEqual(getattr(section, attribute).cm, expected_cm, places=2)

    def test_body_font_is_times_new_roman_fourteen_spaced_one_point_three(self):
        from docx import Document
        from docx.shared import Pt

        path = docx_build.build(self.lesson, self.folder)
        normal = Document(str(path)).styles["Normal"]
        self.assertEqual(normal.font.name, "Times New Roman")
        self.assertEqual(normal.font.size, Pt(14))
        self.assertAlmostEqual(normal.paragraph_format.line_spacing, 1.3, places=3)

    def test_five_twelve_sections_are_present(self):
        path = docx_build.build(self.lesson, self.folder)
        xml = document_xml(path)
        for heading in (
            "I. MỤC TIÊU",
            "II. THIẾT BỊ DẠY HỌC VÀ HỌC LIỆU",
            "III. TIẾN TRÌNH DẠY HỌC",
            "IV. PHỤ LỤC",
        ):
            with self.subTest(heading=heading):
                self.assertIn(heading, xml)

    def test_every_activity_shows_the_four_elements_and_four_steps(self):
        path = docx_build.build(self.lesson, self.folder)
        xml = document_xml(path)
        for label in ("a) Mục tiêu", "b) Nội dung", "c) Sản phẩm", "d) Tổ chức thực hiện"):
            self.assertIn(label, xml)
        for step in (
            "Bước 1. Chuyển giao nhiệm vụ",
            "Bước 2. Học sinh thực hiện nhiệm vụ",
            "Bước 3. Báo cáo, thảo luận",
            "Bước 4. Kết luận, nhận định",
        ):
            self.assertIn(step, xml)

    def test_activities_are_numbered_in_file_order(self):
        path = docx_build.build(self.lesson, self.folder)
        xml = document_xml(path)
        self.assertIn("Hoạt động 1. Mở đầu (10 phút)", xml)
        self.assertIn("Hoạt động 4. Vận dụng (15 phút)", xml)

    def test_activity_title_already_numbered_is_not_doubled(self):
        source = VALID_SOURCE.replace("### Mở đầu", "### Hoạt động 1: Mở đầu")
        lesson = parse.parse_lesson(source)
        path = docx_build.build(lesson, self.folder)
        xml = document_xml(path)
        self.assertIn("Hoạt động 1. Mở đầu (10 phút)", xml)
        self.assertNotIn("Hoạt động 1. Hoạt động 1", xml)

    def test_activity_codes_are_printed(self):
        path = docx_build.build(self.lesson, self.folder)
        xml = document_xml(path)
        self.assertIn("Năng lực số: 1.1.NC1a, 5.1.NC1a", xml)
        self.assertIn("Năng lực AI: AI-H11.2, AI-H11.4", xml)

    def test_worksheet_items_are_separate_left_aligned_paragraphs(self):
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        path = docx_build.build(self.lesson, self.folder)
        document = Document(str(path))
        items = [
            paragraph for paragraph in document.paragraphs
            if paragraph.text.startswith("Câu 1.") or paragraph.text.startswith("Câu 2.")
        ]
        self.assertEqual(len(items), 2, [paragraph.text for paragraph in document.paragraphs])
        for paragraph in items:
            self.assertEqual(paragraph.alignment, WD_ALIGN_PARAGRAPH.LEFT)

    def test_rubric_is_a_four_column_table(self):
        from docx import Document

        path = docx_build.build(self.lesson, self.folder)
        tables = Document(str(path)).tables
        rubric = tables[-1]
        self.assertEqual(len(rubric.columns), 4)
        self.assertEqual(len(rubric.rows), 3)
        self.assertEqual(rubric.rows[0].cells[3].text, "Mức 3")

    def test_chemical_subscripts_are_real_subscripts(self):
        path = docx_build.build(self.lesson, self.folder)
        self.assertIn('w:val="subscript"', document_xml(path))

    def test_internal_notes_never_reach_the_lesson_plan(self):
        source = VALID_SOURCE.replace(
            "- Thầy cô xác nhận giúp em địa chỉ mô phỏng PhET dùng được ở phòng máy của trường.",
            "- DAU-HIEU-GHI-CHU-NOI-BO",
        )
        lesson = parse.parse_lesson(source)
        path = docx_build.build(lesson, self.folder)
        xml = document_xml(path)
        self.assertNotIn("DAU-HIEU-GHI-CHU-NOI-BO", xml)
        self.assertNotIn("Cần thầy cô soát", xml)

    def test_review_file_carries_notes_and_warnings(self):
        source = VALID_SOURCE.replace("thoi-luong: 30", "thoi-luong: 20", 1)
        lesson = parse.parse_lesson(source)
        path = docx_build.write_review(lesson, self.folder, lesson.warnings)
        text = path.read_text(encoding="utf-8")
        self.assertEqual(path.name, "can-soat.md")
        self.assertIn("PhET", text)
        self.assertIn("thời lượng", text)

    def test_review_file_says_so_when_nothing_needs_review(self):
        source = VALID_SOURCE.split("## CAN SOAT")[0]
        lesson = parse.parse_lesson(source)
        path = docx_build.write_review(lesson, self.folder, [])
        self.assertIn("Không có mục nào cần soát", path.read_text(encoding="utf-8"))


class CliTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.folder = Path(self.tmp.name) / "giáo án hoá 11"
        self.folder.mkdir(parents=True)
        self.addCleanup(self.tmp.cleanup)

    def write_source(self, text: str = VALID_SOURCE) -> None:
        (self.folder / "giao-an.md").write_text(text, encoding="utf-8")

    def run_cli(self, *args: str) -> tuple[int, dict]:
        out = io.StringIO()
        with contextlib.redirect_stdout(out), contextlib.redirect_stderr(io.StringIO()):
            code = giao_an.main(list(args))
        printed = out.getvalue().strip().splitlines()
        self.assertEqual(len(printed), 1, printed)
        return code, json.loads(printed[0])

    def test_export_writes_the_plan_and_the_review_file(self):
        self.write_source()
        code, data = self.run_cli("xuat", str(self.folder))
        self.assertEqual(code, 0, data)
        self.assertTrue(data["ready"])
        self.assertEqual([Path(name).name for name in data["files"]], ["giao-an.docx", "can-soat.md"])
        self.assertEqual(data["activities"], 4)
        self.assertEqual(data["periods"], 2)
        self.assertEqual(data["minutes"], 90)
        self.assertEqual(data["nls_codes"], ["1.1.NC1a", "5.1.NC1a"])
        self.assertEqual(data["ai_codes"], ["AI-H11.2", "AI-H11.4"])
        self.assertIsNone(data["error"])

    def test_export_works_with_a_vietnamese_folder_name(self):
        self.write_source()
        code, data = self.run_cli("xuat", str(self.folder))
        self.assertEqual(code, 0, data)
        for name in data["files"]:
            self.assertTrue(Path(name).is_file(), name)

    def test_plan_only_writes_nothing(self):
        self.write_source()
        code, data = self.run_cli("xuat", str(self.folder), "--plan-only")
        self.assertEqual(code, 0, data)
        self.assertEqual(data["files"], [])
        self.assertEqual(list(self.folder.glob("*.docx")), [])

    def test_missing_source_reports_input_step(self):
        code, data = self.run_cli("xuat", str(self.folder))
        self.assertEqual(code, 1)
        self.assertEqual(data["error"]["step"], "input")
        self.assertIn("giao-an.md", data["error"]["message"])

    def test_structure_error_reports_parse_step_with_the_line(self):
        self.write_source(VALID_SOURCE.replace("bao-cao: Hai nhóm trình bày dự đoán.\n", ""))
        code, data = self.run_cli("xuat", str(self.folder))
        self.assertEqual(code, 1)
        self.assertEqual(data["error"]["step"], "parse")
        self.assertIn("Dòng", data["error"]["message"])

    def test_code_error_reports_framework_step(self):
        self.write_source(VALID_SOURCE.replace("1.1.NC1a", "9.9.NC9z"))
        code, data = self.run_cli("xuat", str(self.folder))
        self.assertEqual(code, 1)
        self.assertEqual(data["error"]["step"], "framework")
        self.assertIn("9.9.NC9z", data["error"]["message"])

    def test_missing_python_docx_reports_docx_step(self):
        self.write_source()
        with mock.patch.object(giao_an, "load_docx_build",
                               side_effect=ImportError("No module named 'docx'", name="docx")):
            code, data = self.run_cli("xuat", str(self.folder))
        self.assertEqual(code, 1)
        self.assertEqual(data["error"]["step"], "docx")
        self.assertIn("requirements-vi.txt", data["error"]["fix"])

    def test_unrelated_import_error_reports_internal_step(self):
        self.write_source()
        with mock.patch.object(giao_an, "load_docx_build",
                               side_effect=ImportError("cannot import name 'x'", name="word_parts.base")):
            code, data = self.run_cli("xuat", str(self.folder))
        self.assertEqual(code, 1)
        self.assertEqual(data["error"]["step"], "internal")

    def test_write_failure_reports_write_step(self):
        self.write_source()
        with mock.patch.object(giao_an, "load_docx_build") as loader:
            loader.return_value.build.side_effect = OSError("file đang mở trong Word")
            code, data = self.run_cli("xuat", str(self.folder))
        self.assertEqual(code, 1)
        self.assertEqual(data["error"]["step"], "write")
        self.assertIn("Word", data["error"]["fix"])

    def test_unexpected_failure_still_prints_one_json_line(self):
        self.write_source()
        with mock.patch.object(giao_an, "load_docx_build") as loader:
            loader.return_value.build.side_effect = ValueError("lỗi ngoài dự kiến")
            code, data = self.run_cli("xuat", str(self.folder))
        self.assertEqual(code, 1)
        self.assertEqual(data["error"]["step"], "internal")

    def test_time_mismatch_surfaces_as_a_warning(self):
        self.write_source(VALID_SOURCE.replace("thoi-luong: 30", "thoi-luong: 20", 1))
        code, data = self.run_cli("xuat", str(self.folder))
        self.assertEqual(code, 0, data)
        self.assertTrue(any("thời lượng" in warning for warning in data["warnings"]))

    def test_second_run_warns_about_overwriting(self):
        self.write_source()
        self.run_cli("xuat", str(self.folder))
        code, data = self.run_cli("xuat", str(self.folder))
        self.assertEqual(code, 0, data)
        self.assertTrue(any("Ghi đè" in warning for warning in data["warnings"]))

    def test_cut_textbook_writes_the_slice(self):
        source = self.folder / "sgk.md"
        source.write_text(SGK_SAMPLE, encoding="utf-8")
        code, data = self.run_cli("trich-sgk", str(source), "--bai", "Bài 5")
        self.assertEqual(code, 0, data)
        self.assertTrue(data["heading"].startswith("Bài 5."))
        self.assertEqual(len(data["files"]), 1)
        written = Path(data["files"][0])
        self.assertTrue(written.is_file())
        self.assertIn("đoạn hai", written.read_text(encoding="utf-8"))
        self.assertGreater(data["size_kb"], 0)

    def test_cut_textbook_honours_the_output_path(self):
        source = self.folder / "sgk.md"
        source.write_text(SGK_SAMPLE, encoding="utf-8")
        target = self.folder / "phần bài 5.md"
        code, data = self.run_cli("trich-sgk", str(source), "--bai", "Bài 5", "--ra", str(target))
        self.assertEqual(code, 0, data)
        self.assertEqual(Path(data["files"][0]).name, target.name)

    def test_cut_textbook_missing_file_reports_input_step(self):
        code, data = self.run_cli("trich-sgk", str(self.folder / "khong-co.md"), "--bai", "Bài 5")
        self.assertEqual(code, 1)
        self.assertEqual(data["error"]["step"], "input")

    def test_cut_textbook_unknown_lesson_reports_parse_step(self):
        source = self.folder / "sgk.md"
        source.write_text(SGK_SAMPLE, encoding="utf-8")
        code, data = self.run_cli("trich-sgk", str(source), "--bai", "Bài 99")
        self.assertEqual(code, 1)
        self.assertEqual(data["error"]["step"], "parse")
        self.assertIn("Bài 4.", data["error"]["fix"])

    def test_cut_textbook_refuses_to_overwrite_the_source(self):
        source = self.folder / "sgk.md"
        source.write_text(SGK_SAMPLE, encoding="utf-8")
        code, data = self.run_cli("trich-sgk", str(source), "--bai", "Bài 5", "--ra", str(source))
        self.assertEqual(code, 1)
        self.assertEqual(data["error"]["step"], "input")
        self.assertEqual(source.read_text(encoding="utf-8"), SGK_SAMPLE)

    def test_cut_textbook_write_failure_explains_the_output_path(self):
        source = self.folder / "sgk.md"
        source.write_text(SGK_SAMPLE, encoding="utf-8")
        target = self.folder / "khong-co-thu-muc" / "bai5.md"
        code, data = self.run_cli("trich-sgk", str(source), "--bai", "Bài 5", "--ra", str(target))
        self.assertEqual(code, 1)
        self.assertEqual(data["error"]["step"], "write")
        self.assertIn("--ra", data["error"]["fix"])
        self.assertNotIn("Word", data["error"]["fix"])

    def test_missing_command_reports_input_step(self):
        code, data = self.run_cli()
        self.assertEqual(code, 1)
        self.assertEqual(data["error"]["step"], "input")

    def test_unknown_flag_reports_input_step(self):
        self.write_source()
        code, data = self.run_cli("xuat", str(self.folder), "--khong-co-co-nay")
        self.assertEqual(code, 1)
        self.assertEqual(data["error"]["step"], "input")

    def test_cut_textbook_without_lesson_name_reports_input_step(self):
        source = self.folder / "sgk.md"
        source.write_text(SGK_SAMPLE, encoding="utf-8")
        code, data = self.run_cli("trich-sgk", str(source))
        self.assertEqual(code, 1)
        self.assertEqual(data["error"]["step"], "input")

    def test_broken_stdout_is_not_answered_twice(self):
        self.write_source()

        class BrokenStdout:
            def __init__(self):
                self.writes = 0

            def write(self, text):
                self.writes += 1
                raise BrokenPipeError("bên đọc đã đóng")

            def flush(self):
                pass

        stream = BrokenStdout()
        with mock.patch.object(giao_an.sys, "stdout", stream), contextlib.redirect_stderr(io.StringIO()):
            code = giao_an.main(["xuat", str(self.folder), "--plan-only"])
        self.assertEqual(stream.writes, 1)
        self.assertEqual(code, 0)

    def test_emit_without_buffer_falls_back_to_ascii_json(self):
        class AsciiOnlyStdout:
            def __init__(self):
                self.chunks = []

            def write(self, text):
                text.encode("ascii")
                self.chunks.append(text)
                return len(text)

            def flush(self):
                pass

        stream = AsciiOnlyStdout()
        with mock.patch.object(giao_an.sys, "stdout", stream):
            giao_an.emit({"ready": False, "warnings": ["Thiếu phiếu học tập"]})
        self.assertEqual(len(stream.chunks), 1)
        self.assertEqual(json.loads(stream.chunks[0])["warnings"], ["Thiếu phiếu học tập"])

    def test_emit_falls_back_to_utf8_buffer(self):
        class LegacyStdout:
            def __init__(self):
                self.buffer = io.BytesIO()

            def write(self, text):
                text.encode("cp1252")
                return len(text)

            def flush(self):
                pass

        stream = LegacyStdout()
        with mock.patch.object(giao_an.sys, "stdout", stream):
            giao_an.emit({"ready": False, "warnings": ["Thiếu phiếu học tập"]})
        data = json.loads(stream.buffer.getvalue().decode("utf-8").strip())
        self.assertEqual(data["warnings"], ["Thiếu phiếu học tập"])


if __name__ == "__main__":
    unittest.main()
