"""Dựng ba file Word của đề kiểm tra: đề tiếng Anh, đề song ngữ, đáp án."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

from word_parts import base, inline
from word_parts.base import cell_text as _cell_text
from word_parts.base import clear_borders as _clear_borders
from word_parts.base import fill_cell as _fill_cell
from word_parts.base import grid_borders as _grid_borders
from word_parts.base import set_widths as _set_widths
from word_parts.base import write as _write

from .parse import LEVEL_LABELS, LEVELS, Exam, Question, part_label

GREY = RGBColor(0x59, 0x59, 0x59)
FILENAMES = {"de": "de-en.docx", "song-ngu": "de-song-ngu.docx", "dap-an": "dap-an.docx"}
PART_TITLES = {
    1: "PART I. MULTIPLE CHOICE",
    2: "PART II. TRUE / FALSE",
    3: "PART III. SHORT ANSWER",
}
PART_NOTES = {
    1: "Each question has only one correct answer.",
    2: "For each question, decide whether each statement is true or false.",
    3: "Write your answer in the space provided.",
}
CANDIDATE_LINE = (
    "Full name: ................................................  "
    "Class: ................  Student ID: ................"
)
ANSWER_LINE = "Answer: ............................................"
END_MARKER = "------ THE END ------"
NO_REVIEW = "Không có mục nào cần soát."
SHORT_OPTION = 12
MEDIUM_OPTION = 30
GRID_SIZE = 10


def option_columns(options: dict[str, str]) -> int:
    """Đáp án ngắn xếp 4 cột, trung bình 2 cột, dài thì mỗi đáp án một dòng."""
    longest = max(len(inline.plain_text(text)) for text in options.values())
    if longest <= SHORT_OPTION:
        return 4
    if longest <= MEDIUM_OPTION:
        return 2
    return 1


def _base_document() -> Document:
    """Thể thức đề thi: A4 dọc, lề 1,8/1,8/2,5/1,5 cm, Times New Roman 12, giãn dòng 1,15."""
    return base.new_document()


def _add_header(document: Document, exam: Exam, subtitle: str = "") -> None:
    table = document.add_table(rows=1, cols=2)
    _clear_borders(table)
    _set_widths(table, [Cm(8), Cm(8.5)])
    left, right = table.rows[0].cells
    left_lines = [exam.meta["school"]]
    if exam.meta.get("department"):
        left_lines.append(exam.meta["department"])
    _fill_cell(left, left_lines)
    right_lines = [exam.meta["title"]]
    if subtitle:
        right_lines.append(subtitle)
    right_lines.append(f"Môn: {exam.meta['subject']}")
    right_lines.append(f"Thời gian: {exam.meta['time']} phút")
    if exam.meta.get("code"):
        right_lines.append(f"Mã đề: {exam.meta['code']}")
    _fill_cell(right, right_lines)


def _add_candidate_line(document: Document) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(6)
    _write(paragraph, CANDIDATE_LINE)


def _add_part_heading(document: Document, part: int) -> None:
    heading = document.add_paragraph()
    heading.paragraph_format.space_before = Pt(8)
    _write(heading, PART_TITLES[part], bold=True)
    _write(document.add_paragraph(), PART_NOTES[part], italic=True)


def _add_options(document: Document, options: dict[str, str]) -> None:
    columns = option_columns(options)
    labels = ("A", "B", "C", "D")
    rows = (len(labels) + columns - 1) // columns
    table = document.add_table(rows=rows, cols=columns)
    _clear_borders(table)
    _set_widths(table, [Cm(16.5 / columns)] * columns)
    for index, label in enumerate(labels):
        cell = table.rows[index // columns].cells[index % columns]
        cell.paragraphs[0].text = ""
        _write(cell.paragraphs[0], f"{label}. ", bold=True)
        _write(cell.paragraphs[0], options[label])


def _add_statements(document: Document, question: Question) -> None:
    table = document.add_table(rows=len(question.statements) + 1, cols=3)
    _grid_borders(table)
    _set_widths(table, [Cm(12.5), Cm(2), Cm(2)])
    for cell, text in zip(table.rows[0].cells, ("Statement", "True", "False")):
        _cell_text(cell, text, bold=True)
    for row_index, (label, text, _) in enumerate(question.statements, start=1):
        cells = table.rows[row_index].cells
        cells[0].paragraphs[0].text = ""
        _write(cells[0].paragraphs[0], f"{label}) ")
        _write(cells[0].paragraphs[0], text)


def _add_question(document: Document, question: Question, *, bilingual: bool) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4)
    _write(paragraph, f"Question {question.number}. ", bold=True)
    _write(paragraph, question.en)
    if bilingual:
        translation = document.add_paragraph()
        _write(translation, question.vi or "(chưa có bản tiếng Việt)", italic=True, color=GREY)
    if question.part == 1:
        _add_options(document, question.options)
    elif question.part == 2:
        _add_statements(document, question)
    else:
        _write(document.add_paragraph(), ANSWER_LINE)


def _add_end(document: Document) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(8)
    _write(paragraph, END_MARKER, bold=True)


def build_de(exam: Exam, path: Path, *, bilingual: bool = False) -> Path:
    document = _base_document()
    _add_header(document, exam, subtitle="BẢN SONG NGỮ" if bilingual else "")
    _add_candidate_line(document)
    for part in (1, 2, 3):
        questions = exam.part(part)
        if not questions:
            continue
        _add_part_heading(document, part)
        for question in questions:
            _add_question(document, question, bilingual=bilingual)
    _add_end(document)
    document.save(str(path))
    return path


def _add_answer_grid(document: Document, exam: Exam) -> None:
    part1 = exam.part(1)
    if part1:
        _write(document.add_paragraph(), PART_TITLES[1], bold=True)
        for start in range(0, len(part1), GRID_SIZE):
            chunk = part1[start:start + GRID_SIZE]
            table = document.add_table(rows=2, cols=len(chunk))
            _grid_borders(table)
            for index, question in enumerate(chunk):
                _cell_text(table.rows[0].cells[index], str(question.number), bold=True)
                _cell_text(table.rows[1].cells[index], question.key)
    part2 = exam.part(2)
    if part2:
        _write(document.add_paragraph(), PART_TITLES[2], bold=True)
        table = document.add_table(rows=len(part2) + 1, cols=5)
        _grid_borders(table)
        for cell, text in zip(table.rows[0].cells, ("Question", "a", "b", "c", "d")):
            _cell_text(cell, text, bold=True)
        for row_index, question in enumerate(part2, start=1):
            cells = table.rows[row_index].cells
            _cell_text(cells[0], str(question.number))
            for column, (_, _, correct) in enumerate(question.statements, start=1):
                _cell_text(cells[column], "T" if correct else "F")
    part3 = exam.part(3)
    if part3:
        _write(document.add_paragraph(), PART_TITLES[3], bold=True)
        table = document.add_table(rows=len(part3) + 1, cols=3)
        _grid_borders(table)
        for cell, text in zip(table.rows[0].cells, ("Question", "Answer", "Unit")):
            _cell_text(cell, text, bold=True)
        for row_index, question in enumerate(part3, start=1):
            cells = table.rows[row_index].cells
            for cell, text in zip(cells, (str(question.number), question.key, question.unit or "—")):
                _cell_text(cell, text)


def _add_scoring(document: Document, exam: Exam) -> None:
    heading = document.add_paragraph()
    heading.paragraph_format.space_before = Pt(8)
    _write(heading, "Thang điểm", bold=True)
    counts = exam.counts()
    lines = []
    if counts["part1"]:
        lines.append(f"Phần I: {counts['part1']} câu × 0,25 = {counts['part1'] * 0.25:g} điểm.")
    if counts["part2"]:
        lines.append(
            f"Phần II: {counts['part2']} câu × 1,0 = {counts['part2'] * 1.0:g} điểm "
            "(đúng 1 ý 0,1; 2 ý 0,25; 3 ý 0,5; 4 ý 1,0)."
        )
    if counts["part3"]:
        lines.append(f"Phần III: {counts['part3']} câu × 0,25 = {counts['part3'] * 0.25:g} điểm.")
    lines.append(f"Tổng: {exam.total_points():g} điểm.")
    for line in lines:
        _write(document.add_paragraph(), line)


def _add_explanations(document: Document, exam: Exam) -> None:
    explained = [question for question in exam.questions if question.why]
    if not explained:
        return
    heading = document.add_paragraph()
    heading.paragraph_format.space_before = Pt(8)
    _write(heading, "Hướng dẫn giải", bold=True)
    for question in explained:
        paragraph = document.add_paragraph()
        _write(paragraph, f"PART {part_label(question.part)} — Question {question.number}. ", bold=True)
        _write(paragraph, question.why)


def _add_matrix(document: Document, exam: Exam) -> None:
    heading = document.add_paragraph()
    heading.paragraph_format.space_before = Pt(8)
    _write(heading, "Ma trận đặc tả", bold=True)
    rows = exam.matrix()
    table = document.add_table(rows=len(rows) + 2, cols=5)
    _grid_borders(table)
    _set_widths(table, [Cm(8.5), Cm(2), Cm(2), Cm(2), Cm(2)])
    header = ("Chủ đề", *(LEVEL_LABELS[level] for level in LEVELS), "Tổng")
    for cell, text in zip(table.rows[0].cells, header):
        _cell_text(cell, text, bold=True)
    for row_index, row in enumerate(rows, start=1):
        for cell, value in zip(table.rows[row_index].cells, row):
            _cell_text(cell, str(value))
    totals = (
        "Tổng",
        *(str(sum(row[index] for row in rows)) for index in (1, 2, 3)),
        str(sum(row[4] for row in rows)),
    )
    for cell, value in zip(table.rows[-1].cells, totals):
        _cell_text(cell, value, bold=True)


def _add_review(document: Document, exam: Exam, warnings: list[str]) -> None:
    heading = document.add_paragraph()
    heading.paragraph_format.space_before = Pt(8)
    _write(heading, "Cần thầy cô soát", bold=True)
    items = list(exam.review_notes) + list(warnings)
    if not items:
        _write(document.add_paragraph(), NO_REVIEW)
        return
    for item in items:
        _write(document.add_paragraph(), f"- {item}")


def build_dap_an(exam: Exam, path: Path, warnings: list[str]) -> Path:
    document = _base_document()
    _add_header(document, exam, subtitle="ĐÁP ÁN VÀ HƯỚNG DẪN CHẤM")
    _add_answer_grid(document, exam)
    _add_scoring(document, exam)
    _add_explanations(document, exam)
    _add_matrix(document, exam)
    _add_review(document, exam, warnings)
    document.save(str(path))
    return path


def build(exam: Exam, folder: Path, parts: list[str], warnings: list[str]) -> list[Path]:
    written: list[Path] = []
    for part in parts:
        path = folder / FILENAMES[part]
        if part == "de":
            build_de(exam, path)
        elif part == "song-ngu":
            build_de(exam, path, bilingual=True)
        else:
            build_dap_an(exam, path, warnings)
        written.append(path)
    return written
